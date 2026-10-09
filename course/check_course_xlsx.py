# -*- coding: utf-8 -*-
"""表計算ブック（教員版・学習者版）の検査。

    python course/check_course_xlsx.py              # 答え語の検査＋構造の検査（Excel なしで動く）
    python course/check_course_xlsx.py --pandas     # ＋ pandas による式の再現（4バンド・密度比・USGS年代・
                                                    #    上位10の同点解消の網羅）
    python course/check_course_xlsx.py --excel      # ＋ Excel 実機（COM）での再計算・入力規則・上位10の網羅・
                                                    #    状態表示・範囲表示の検査（Windows＋Excel が必要）
    python course/check_course_xlsx.py --all        # すべて

1. 答え語の検査（要件 §4-X-2）
   学習者版の「全セルの文字列」（非表示シートを含む。数式の文字列も含む）と、xlsx 内のすべての XML
   （共有文字列・チャート・入力規則の文言・名前・文書情報）から、答え語リスト ANSWER_WORDS の出現を探す。
   学習者版で 0 件であること、**教員版にはリストの全語が存在する**こと（検査が空振りでないこと）を確認する。
2. 構造の検査（要件 §3.2）
   シート名と順序が両版で同じ／学習者版で非表示にしたシートの一覧／黄色い入力セルの位置と既定値／
   C5 のプルダウンと重みの入力規則／学習者版の データ_地域 が A〜E 列だけ、など。
3. 値の検査（--pandas / --excel）
   同じ式を pandas で再現し、Excel 実機の再計算結果と突き合わせる。
   第2時の再設計（docs/design_dai2ji_v2.md rev2）：4バンド（295・279・234・156 K）、データ_緯度行（90行）、
   データ_地点比較（5クレーター）、ステップ1b の線の式、ステップ1c の表・カーブを、同梱の生データ
   （data/diviner_global.csv.gz・data/craters_subset.csv）から pandas で独立に再現して突き合わせる。

4. 要素ごとの分析と重み（docs/design_yoso_v3.md rev2）
   ステップ4の C5 のプルダウン（先頭＝『月全体（線の内側）』。非表示列 M5:M13・名前 AreaChoices／AllMoon／LineOK）と C4（線）、
   学習者版の既定（C8＝1）、データ_環境 O 列の『採点する行』、A6・B36 の月全体の分岐、H 列（夜の最低温度）の構造・式の検査。
   月全体の期待値（設定A〜D・電力0の試行(b)・線50〜90・上位10・範囲・状態表示・B36 の確認の一文、4b の水1番・傾斜2番）を
   pandas で再現し（--pandas）、Excel 実機の再計算結果と突き合わせる（--excel）。負の検査（C4 の空欄・範囲外、電力の行だけ、重みが全部0、C8＝0）も含む。
   学習者版に『ミッション』『観測』『合格』『目安』『ティコ付近』『表側の中央』と答えの座標の文がないこと（AC1④⑤）も走査する。

終了コード: すべて合格なら 0、1つでも不合格なら 1。
元のファイルは Excel で開かない（コピーを一時フォルダに作って開く）。
"""
import argparse
import json
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE = pathlib.Path(__file__).resolve().parent
DATA = HERE / "data"
TEACHER = HERE / "course_moonbase.xlsx"
STUDENT = HERE / "student" / "course_moonbase_student.xlsx"
DRIVER = HERE / "excel_driver.ps1"
ROOT_DATA = HERE.parent / "data"          # 同梱の生データ（diviner_global.csv.gz・craters_subset.csv）

# ---------------------------------------------------------------------------
# 答え語のリスト。教員版の実物のセルから拾った「答えにあたる文・語」と、要件 §4-X-2 の語。
# 学習者版では 0 件、教員版では全語が 1 件以上あること。
# ---------------------------------------------------------------------------
ANSWER_WORDS = [
    # 要件 §4-X-2 で挙げられた語のうち、教員版に残るもの
    "第一候補", "電波天文に理想", "永久影のそば",
    # ステップ1
    "極付近（要注意）",
    # ステップ2（企画書 G1 の A33、G3 の D26）
    "海岸線の外側の陸も丸に入る", "落ちた後に消えた", "隕石が落ちなかった", "C の年代の平均が小さいほう",
    # ステップ3（企画書 G1）
    "正＝表側で通信できる", "負＝裏側で電波が静か", "裏側は地球の仰角が負", "南極は日較差が小さいが太陽高度は",
    # ステップ4（教員版の重みの行の注。用途の語は学習者版に出さない）
    "（電力）", "（温度の安定）", "（通信できる）", "（観測用。授業では使わない）", "（夜に冷えすぎない）",
    # ステップ5
    "Artemis III（有人・氷）",
    # データ_地域 の rationale / caveat
    "永久影に氷があり", "電波雑音が届かない", "電波静穏度が最も高い", "LCRT（月裏側電波望遠鏡）構想の対象域",
    # 参考
    "（＝陸のほうが古い）",
    # 要素ごとの分析と重み（設計 docs/design_yoso_v3.md）で教員版に入れた答え・組合せ例の文（学習者版には出さない）
    "月全体の例", "水を1番にした班の例", "試行(a) 1番を替える", "電力を0にすると", "北と南で同点",
    "選んだのは重みではなく線", "重みが違う班どうしで比べない", "傾斜（建設）は2番", "1位は行の順で決まっただけ",
    "使うシートは1番で決まります", "1番に置いた要素が違うと",
]

# 役目を終えた旧い教員版の答え語（ミッション版の文）。教員版には無いが、学習者版に紛れ込んでいないことを確かめ続ける（学習者版0件）
RETIRED_ANSWER_WORDS = [
    "氷採掘は南極", "見えたら失格", "裏側に", "電波天文台は裏側に", "通信重視なら表側に", "半球ごと変わります",
    "（＝月の裏側）", "氷がありそうなこと", "ずっと日が当たらない永久影", "基地に必要なこと", "各ミッションの答え",
    "（発電）", "（熱の安定）", "（電波が静か・天文台）", "経度180度あたりが上位に来る", "ミッションごとの選び方",
    "2・2・0・0・1", "0・2・0・3・0", "1・1・3・0・0", "1・2・1・0・1", "永久影は環境データに無い",
    "赤道の海（静かの海） か 南極", "氷採掘・南極を選んだ班向け", "ミッションごとの重み", "氷採掘基地", "太陽光発電基地",
    "有人基地", "氷採掘は南極で合意", "実在の計画は、ミッションで半球がちがう",
]

# 第2時の再設計（設計 §7.1）で足した語。学習者版では 0 件。教員版は、教員版に入れた語（TEACHER_NEW_PRESENT）だけ存在を確認する。
# 「作り方」だけでは、既存の案内（「作り方は course/build_course_data.py」など）に当たるので、「データの作り方」にしてある。
NEW_ANSWER_WORDS = [
    "平らな地面", "斜面", "データの作り方", "本当の姿", "分かっていません", "データの限界",
    "信頼できない", "信頼できる", "信頼しにくい", "信じにくい", "信じてよい", "岩が多い", "岩塊", "冷めにくい",
]
TEACHER_NEW_PRESENT = ["平らな地面", "斜面", "データの作り方"]
ALL_ANSWER_WORDS = ANSWER_WORDS + RETIRED_ANSWER_WORDS + NEW_ANSWER_WORDS
# 新しい2つのデータ表の見出しに入れてはいけない語（判断・原因の語）
HEADER_BAD_WORDS = ["信頼", "信じ", "判定", "外れ", "違う", "異常", "原因", "斜面", "岩", "限界", "要注意"]

STUDENT_HIDDEN = ["ステップ3_月全体", "ステップ4b_南極", "ステップ5_まとめ",
                  "データ_北極", "データ_地質", "データ_クレーター年代"]
S1B, S1C = "ステップ1b_帯を刻む", "ステップ1c_地点と帯"
SHEETS = ["はじめに", "ステップ1_温度", S1B, S1C, "ステップ2_海と陸", "ステップ3_月全体", "ステップ4_地域を選ぶ",
          "ステップ4b_南極", "ステップ4b_スコア", "ステップ5_まとめ", "データ_温度", "データ_緯度行", "データ_地点比較", "データ_クレーター",
          "データ_クレーター年代", "データ_環境", "データ_地域", "データ_南極", "データ_北極",
          "データ_地質", "データ_着陸地点", "参考"]
YELLOW = "FFF6E9"   # 黄色い入力セルの色（アルファ部分は openpyxl 出力＝00、Excel 保存＝FF と違うので比べない）
REG_FAR_EQ = "裏側・赤道（月の裏側の赤道帯）"
ALL_MOON = "月全体（線の内側）"      # ステップ4 C5 のプルダウンの先頭（非表示列 M5）
S1, S2, S4, S4B = "ステップ1_温度", "ステップ2_海と陸", "ステップ4_地域を選ぶ", "ステップ4b_スコア"

RESULTS = []   # (合否, 項目, 詳細)


def record(ok, name, detail=""):
    RESULTS.append((bool(ok), name, detail))
    print(("[OK]   " if ok else "[NG]   ") + name + (f"  {detail}" if detail else ""))


# ---------------------------------------------------------------------------
# 1. 答え語の検査
# ---------------------------------------------------------------------------
def scan_cells(path):
    """全シート（非表示を含む）の全セルの文字列（数式は式の文字列）→ {語: [場所,…]}"""
    wb = load_workbook(path, read_only=True)       # read_only でも非表示シートを含めて全部読める
    texts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value:
                    texts.append((ws.title, c.coordinate, c.value))
    hits = {}
    for w in ALL_ANSWER_WORDS:
        for sh, cell, t in texts:
            if w in t:
                hits.setdefault(w, []).append(f"{sh}!{cell}")
    return hits, len(texts)


def scan_xml(path):
    """xlsx の全部品（共有文字列・シート・チャート・名前・入力規則・文書情報）を文字列として探す"""
    hits = {}
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if not n.endswith((".xml", ".rels")):
                continue
            t = z.read(n).decode("utf-8", errors="replace")
            for w in ALL_ANSWER_WORDS:
                if w in t:
                    hits.setdefault(w, []).append(n)
    return hits


def check_answers(teacher, student):
    print("\n== 1. 答え語の検査 ==")
    print(f"答え語リスト: {len(ANSWER_WORDS)} 語＋第2時で追加 {len(NEW_ANSWER_WORDS)} 語")
    hs, n_s = scan_cells(student)
    hx = scan_xml(student)
    record(not hs, f"学習者版：全セル（{n_s} 個の文字列セル、非表示シート含む）に答え語 0 件",
           "" if not hs else "; ".join(f"「{w}」→{v[:3]}" for w, v in hs.items()))
    record(not hx, "学習者版：xlsx 内の全 XML（共有文字列・チャート・入力規則・名前・文書情報）に答え語 0 件",
           "" if not hx else "; ".join(f"「{w}」→{v[:3]}" for w, v in hx.items()))
    ht, n_t = scan_cells(teacher)
    need_t = ANSWER_WORDS + TEACHER_NEW_PRESENT
    missing = [w for w in need_t if w not in ht]
    record(not missing, f"教員版：リストの全 {len(ANSWER_WORDS)} 語と、第2時の教員用メモの語 {len(TEACHER_NEW_PRESENT)} 語が存在（検査が空振りでない。{n_t} 個の文字列セル中）",
           "" if not missing else "教員版に無い語: " + "、".join(missing))
    # 新しい4シートが走査対象に入っていること（文字列セルがある）。学習者版の新シートに答え語が 0 件
    wbs = load_workbook(student, read_only=True)
    cnt = {}
    for nm in (S1B, S1C, "データ_緯度行", "データ_地点比較"):
        cnt[nm] = sum(1 for row in wbs[nm].iter_rows() for c in row if isinstance(c.value, str) and c.value)
    record(all(v > 0 for v in cnt.values()) and not any(
        h.split("!")[0] in cnt for v in hs.values() for h in v),
        "学習者版：新しい4シート（1b・1c・データ_緯度行・データ_地点比較）も走査対象で、答え語 0 件",
        "文字列セル数 " + "、".join(f"{k} {v}" for k, v in cnt.items()))
    # 空振りでないことの追加確認：検査関数に教員版を食わせて検出数を出す
    print(f"       教員版での検出：{sum(len(v) for v in ht.values())} 箇所／{len(ht)} 語")
    # 非表示シートを実際に走査していること（学習者版の非表示シートの文字列数を数える）
    wbs = load_workbook(student, read_only=True)
    hidden_cnt = sum(1 for ws in wbs.worksheets if ws.sheet_state != "visible")
    record(hidden_cnt == len(STUDENT_HIDDEN), f"学習者版の非表示シート数 {hidden_cnt}（走査対象に含まれる）")


# ミッション名・用途を示す語（要件：地域・半球とミッションの対応を学習者版に書かない）。
# 学習者版の全セル（非表示シート・数式の文字列を含む）と、xlsx 内の全部品（定義名・チャート・入力規則の文言・
# 文書情報など）に出たら不合格。出た場合は場所と前後の文脈を表示する。
MISSION_WORDS = ["電波", "氷採掘", "天文", "発電", "通信", "有人", "採掘", "太陽光"]
# 許容リスト（最小限）。「地域名・データ列の値として出る語」で、ミッションとの対応を示さないもの。
# 形式：(場所の前方一致, その語を含む文脈の部分文字列, 理由)。許容は1件ずつ理由を書く。
MISSION_ALLOWED = [
    ("データ_着陸地点!B", "着陸(有人)", "実在の着陸地点の種別（有人／無人）。史実の属性で、ミッションと地域の対応ではない"),
    ("データ_着陸地点!B", "着陸候補(有人)", "同上（Artemis の着陸候補地の種別）"),
]


def _ctx(text, i, n):
    return text[max(0, i - 14): i + n + 14].replace(chr(10), " ")


def mission_hits(path):
    """[(場所, 語, 文脈)] を返す（許容リストに当たるものは除く）。セルと、セル以外の XML 部品の両方を見る"""
    hits = []
    wb = load_workbook(path, read_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    for w in MISSION_WORDS:
                        i = c.value.find(w)
                        while i >= 0:
                            hits.append((f"{ws.title}!{c.coordinate}", w, _ctx(c.value, i, len(w))))
                            i = c.value.find(w, i + 1)
    cell_ctx = {(w, ctx) for _, w, ctx in hits}
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if not n.endswith((".xml", ".rels")) or n == "xl/sharedStrings.xml":
                continue
            t = z.read(n).decode("utf-8", errors="replace")
            for w in MISSION_WORDS:
                i = t.find(w)
                while i >= 0:
                    ctx = _ctx(t, i, len(w))
                    # セルの値（数式を除く）はセルの走査で拾い済み。ここでは重複を省く
                    if (w, ctx) not in cell_ctx:
                        hits.append((n, w, ctx))
                    i = t.find(w, i + 1)
    out, allowed = [], []
    for where, w, ctx in hits:
        if any(where.startswith(a) and sub in ctx for a, sub, _ in MISSION_ALLOWED):
            allowed.append((where, w, ctx))
            continue
        out.append((where, w, ctx))
    return out, allowed


def check_mission_words(teacher, student):
    print(chr(10) + "== 1b. ミッション語の検査 ==")
    print("語：" + "・".join(MISSION_WORDS) + f"　許容リスト {len(MISSION_ALLOWED)} 件")
    hs, al = mission_hits(student)
    if al:
        print(f"       許容した出現 {len(al)} 件：" + "、".join(sorted({c for _, _, c in al})) + "（データ_着陸地点 の種別列）")
    record(not hs, "学習者版：全セル・全部品にミッション語 0 件",
           "" if not hs else f"{len(hs)} 件。" + " | ".join(f"{w}@{a}『{c}』" for a, w, c in hs[:8]))
    for a, w, c in hs[:30]:
        print(f"       学習者版 {a}  「{w}」 …{c}…")
    ht, _ = mission_hits(teacher)
    found = {w for _, w, _ in ht}
    # 「ミッション」を「要素」に置き換えた後の教員版には、データ_地域の rationale・実在の計画の表・要素の表に残る語だけがある
    # （発電・太陽光は旧ミッション名の語で、教員版からも消えた）。検査が空振りでないこと＝残る6語が検出されること
    want = set(MISSION_WORDS) - {"発電", "太陽光"}
    record(found == want, "教員版：残る 6 語（電波・氷採掘・天文・採掘・通信・有人）が検出される（検査が空振りでない。発電・太陽光は旧ミッション名の語で教員版からも消えた）",
           f"検出 {len(ht)} 件／語 {sorted(found)}" + ("" if found == want else f" 期待との差 {sorted(found ^ want)}"))


# 要素ごとの分析と重み（設計 docs/design_yoso_v3.md rev2 AC1④⑤）：学習者版に、旧い呼び名（ミッション）・発展の語（観測）・
# 評価の語（合格・目安）と、答えにあたる記述（月全体の1位の座標・「ティコ付近」「表側の中央」）がない
V3_STUDENT_WORDS = ["ミッション", "観測", "合格", "目安", "ティコ付近", "表側の中央"]
# 学習者版に載せてはいけない座標（月全体の1位・4bの1位）。(lat, lon) を文として書いたものを探す
ANSWER_POINTS = [(-43.5, -11.5), (-1.5, 0.5), (1.5, 0.5), (-4.5, -8.5), (88.5, -103.5), (-85.6, 138.0), (-84.4, 156.0),
                 (-82.5, 10.5), (73.5, -11.5), (-88.4, -147.0)]


def _coord_regex(lat, lon):
    def num(x):
        t = f"{abs(x):.1f}"
        return (r"[-−‐]\s*" if x < 0 else "") + re.escape(t)
    return re.compile(num(lat) + r"\s*[,，、]\s*" + num(lon))


def check_v3_student_words(teacher, student):
    print(chr(10) + "== 1c. 旧い呼び名・発展の語・答えの座標（要素ごとの分析と重み。AC1④⑤・AC2） ==")
    wb = load_workbook(student, read_only=True)
    texts = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and c.value]
    hits = [f"{sh}!{a}「{w}」" for sh, a, t in texts for w in V3_STUDENT_WORDS if w in t]
    # 「ティコ」だけ（クレーター名。ステップ1c のプルダウン・データ表）は許容。ここでは「ティコ付近」だけを禁止している
    xml_hits = []
    with zipfile.ZipFile(student) as z:
        for n in z.namelist():
            if n.endswith((".xml", ".rels")):
                t = z.read(n).decode("utf-8", errors="replace")
                xml_hits += [f"{n}「{w}」" for w in V3_STUDENT_WORDS if w in t]
    record(not hits and not xml_hits, f"学習者版：全セルと全部品に {V3_STUDENT_WORDS} が 0 件（クレーター名の『ティコ』は許容）",
           "; ".join((hits + xml_hits)[:6]))
    pat = [(p_, _coord_regex(*p_)) for p_ in ANSWER_POINTS]
    chits = [f"{sh}!{a} {p_}" for sh, a, t in texts for p_, rx in pat if rx.search(t)]
    record(not chits, "学習者版：月全体・4bの1位の座標を文として書いた箇所が 0 件（文字列セルを走査。数値セルのデータは対象外）", "; ".join(chits[:5]))
    # 空振りの確認：教員版の月全体の例・4bの例には座標の文字列がある
    wt = load_workbook(teacher, read_only=True)
    tt = [c.value for ws in wt.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    found = [p_ for p_, rx in pat if any(rx.search(t) for t in tt)]
    record(len(found) >= 3, f"（空振りでない確認）教員版には座標を文にした箇所がある：{found[:4]}")
    # 旧い呼び名の「ミッション」が教員版の表計算の文にも残っていない（実在の計画の説明を除き、0 件）
    tm = [t for t in tt if "ミッション" in t]
    record(not tm, "教員版：旧い呼び名『ミッション』が 0 件", "; ".join(t[:20] for t in tm[:3]))


# 地域名（データ_地域 A列）の列挙と、ミッション語・答えの示唆がないことの確認
def check_region_names(path, label):
    wb = load_workbook(path, read_only=True)
    ws = wb["データ_地域"]
    names = [r[0].value for r in ws.iter_rows(min_row=2, max_row=9, max_col=1)]
    bad = [n for n in names if any(w in n for w in MISSION_WORDS)]
    record(not bad and len(names) == 8, f"{label}：地域名 8 件にミッション語がない", " / ".join(names))


PII_PATTERNS = ["C:\\Users", "OneDrive", "@", "lastModifiedBy", "absPath"] + [os.environ.get("USERNAME", "")]
PII_PATTERNS = [w for w in PII_PATTERNS if w]


def check_pii(path, label):
    """xlsx 内の全部品に、利用者名・ローカルパス・メールの断片、Excel が埋める最終更新者・絶対パスがないこと"""
    found = {}
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            t = z.read(n).decode("utf-8", errors="replace")
            for w in PII_PATTERNS:
                if w in t:
                    found.setdefault(w, []).append(n)
    record(not found, f"{label}：個人名・ローカルパス・最終更新者の記録が xlsx 内にない",
           "" if not found else "; ".join(f"{w}→{v[:2]}" for w, v in found.items()))


# ---------------------------------------------------------------------------
# 2. 構造の検査
# ---------------------------------------------------------------------------
def fill_of(ws, addr):
    c = ws[addr]
    return str(c.fill.fgColor.rgb)[-6:] if c.fill and c.fill.fill_type == "solid" else None


def check_structure(teacher, student):
    print("\n== 2. 構造の検査 ==")
    wt = load_workbook(teacher)
    ws_ = load_workbook(student)
    record(wt.sheetnames == SHEETS and ws_.sheetnames == SHEETS,
           "シート名・順序が両版で同じ（既存の18シート＋第2時の新しい4シート＝22シート）")
    record(all(s.sheet_state == "visible" for s in wt.worksheets), "教員版：非表示シートなし")
    hid = [s.title for s in ws_.worksheets if s.sheet_state != "visible"]
    record(sorted(hid) == sorted(STUDENT_HIDDEN), "学習者版の非表示シート", "、".join(hid))

    # 黄色い入力セル（要件 §3.2：位置を動かさない）
    inputs = {S4: ["C4", "C5", "C8", "C9", "C10", "C11", "C12"], S4B: ["C8", "C9", "C10", "C11"],
              S1: ["A8", "B8", "A9", "B9", "A10", "B10", "A11", "B11", "B17", "B18"],
              S1B: ["B6", "B7"], S1C: ["B5"]}
    for label, wb in (("教員版", wt), ("学習者版", ws_)):
        bad = [f"{sh}!{a}" for sh, lst in inputs.items() for a in lst if fill_of(wb[sh], a) != YELLOW]
        record(not bad, f"{label}：黄色い入力セルの位置（ステップ4 C4〔線。新設〕・C5・C8〜C12、4b C8〜C11、ステップ1、1b B6・B7、1c B5）", "、".join(bad))
    # 既定値
    t4, t4b = wt[S4], wt[S4B]
    s4, s4b = ws_[S4], ws_[S4B]
    record(t4["C5"].value == "赤道の海（静かの海）" and [t4[f"C{r}"].value for r in range(8, 13)] == [2, 2, 0, 0, 0]
           and [t4b[f"C{r}"].value for r in range(8, 12)] == [3, 0, 0, 2],
           "教員版：既定値（C5＝赤道の海、重み 2・2・0・0・0／4b 3・0・0・2）は従来どおり")
    record(t4["C4"].value == 70, "教員版：ステップ4 C4（線）の既定は 70")
    # 更新：学習者版の既定の重みは、C8（電力＝太陽高度の行／4bは日照率の行）＝1 で配る（設計 V3-S5。以前は重みがすべて 0 だった）
    record(s4["C5"].value in (None, "") and s4["C4"].value in (None, "") and [s4[f"C{r}"].value for r in range(8, 13)] == [1, 0, 0, 0, 0]
           and [s4b[f"C{r}"].value for r in range(8, 12)] == [1, 0, 0, 0],
           "学習者版：C5・C4（線）は空欄、重みは C8（電力）＝1・ほかは 0（ステップ4・4b。更新：以前はすべて 0）")

    # 入力規則
    for label, wb in (("教員版", wt), ("学習者版", ws_)):
        dvs = wb[S4].data_validations.dataValidation
        lst = [d for d in dvs if d.type == "list" and "C5" in str(d.sqref)]
        line_dv = [d for d in dvs if d.type == "whole" and "C4" in str(d.sqref) and d.formula1 == "50" and d.formula2 == "90"]
        whole = [d for d in dvs if d.type == "whole" and "C8:C12" in str(d.sqref)
                 and d.formula1 == "0" and d.formula2 == "5"]
        dvs_b = wb[S4B].data_validations.dataValidation
        whole_b = [d for d in dvs_b if d.type == "whole" and "C8:C11" in str(d.sqref)
                   and d.formula1 == "0" and d.formula2 == "5"]
        ref = wb.defined_names.get("RegionNames")
        ac, am, lo = (wb.defined_names.get(k) for k in ("AreaChoices", "AllMoon", "LineOK"))
        record(len(lst) == 1 and lst[0].formula1 == "AreaChoices" and ref is not None
               and ref.attr_text == "データ_地域!$A$2:$A$9"
               and ac is not None and ac.attr_text == f"{S4}!$M$5:$M$13"
               and am is not None and am.attr_text == f"{S4}!$M$5" and lo is not None and lo.attr_text == f"{S4}!$M$14",
               f"{label}：C5 のプルダウン（名前 AreaChoices＝{S4}!M5:M13、AllMoon＝M5、LineOK＝M14。RegionNames＝データ_地域!A2:A9 は残る）")
        s4w = wb[S4]
        record(s4w.column_dimensions["M"].hidden and s4w["M5"].value == ALL_MOON
               and [s4w[f"M{6 + i}"].value for i in range(8)] == [f"=データ_地域!A{2 + i}" for i in range(8)]
               and str(s4w["M14"].value).startswith("=IFERROR(IF(AND(ISNUMBER($C$4)"),
               f"{label}：非表示列 M：M5＝『{ALL_MOON}』（プルダウンの先頭）、M6:M13＝データ_地域 A2:A9 へのリンク、M14＝LineOK の式")
        record(len(line_dv) == 1 and str(s4w["A4"].value).startswith("線 [°]") and "A4:B4" in [str(r) for r in s4w.merged_cells.ranges]
               and "温度のデータを使う範囲" in str(s4w["A4"].value) and "採点する" not in str(s4w["A4"].value)
               and "温度のデータを使う範囲" in str(line_dv[0].prompt) and "採点" not in str(line_dv[0].prompt),
               f"{label}：ステップ4 C4（線）の入力規則（整数 50〜90。吹き出しは『温度のデータを使う範囲』）と A4 のラベル")
        record(len(whole) == 1 and len(whole_b) == 1,
               f"{label}：重みの入力規則（ステップ4 C8:C12、4b C8:C11＝0〜5 の整数）")
    # 状態表示・範囲表示セルの存在
    for label, wb in (("教員版", wt), ("学習者版", ws_)):
        a6 = str(wb[S4]["A6"].value)
        a13 = str(wb[S4B]["A13"].value)
        b36 = str(wb[S4]["B36"].value)
        b36b = str(wb[S4B]["B36"].value)
        record(a6.startswith("=") and "南極" in a6 and a13.startswith("=") and b36.startswith("=")
               and "日付変更線" in b36 and b36b.startswith("=") and "日照率" in b36b
               and "AllMoon" in a6 and "LineOK" in a6 and "AllMoon" in b36 and "LineOK" in b36,
               f"{label}：状態表示（ステップ4 A6・4b A13）と範囲表示（B36）の式がある")
    # 順位用列（同点解消）
    for label, wb in (("教員版", wt), ("学習者版", ws_)):
        e = wb["データ_環境"]
        p = wb["データ_南極"]
        o2 = str(e["O2"].value)
        record("AllMoon" in o2 and "LineOK" in o2 and "ABS($A2)<=LineOK" in o2 and "ISNUMBER(LineOK)" in o2,
               f"{label}：データ_環境 O 列の『採点する行』の条件（月全体＝|緯度|≦線、地域名＝region 一致）")
        record(str(e["P2"].value).startswith("=IF(O2") and "ROW()" in str(e["P2"].value)
               and "ROW()" in str(p["M2"].value) and e["O1"].value == "スコア" and p["L1"].value == "スコア",
               f"{label}：順位用列（データ_環境 P、データ_南極 M）に行番号のタイブレークがある")
        f = str(wb[S4]["B26"].value)
        record("LARGE(データ_環境!$P$2" in f and "INDEX(データ_環境!$O$2" in f,
               f"{label}：上位10は順位用列で引き、表示のスコアは元のスコア列（見た目は変わらない）")
    # 学習者版の データ_地域
    dr = ws_["データ_地域"]
    record(dr.max_column == 5 and dr["A1"].value == "name" and dr["E1"].value == "lon_max",
           "学習者版：データ_地域は A〜E 列だけ（rationale・caveat なし）", f"max_column={dr.max_column}")
    record(wt["データ_地域"].max_column == 7, "教員版：データ_地域は A〜G 列（rationale・caveat あり）")
    # 地域名とプルダウン一覧の一致
    names = [dr[f"A{r}"].value for r in range(2, 10)]
    env_regions = set(pd.read_csv(DATA / "env_grid.csv")["region"].dropna())
    record(set(names) == env_regions, "データ_地域 A2:A9 の名前が データ_環境 の region 列の値と一致", f"{len(names)} 件")
    # 学習者版の「はじめに」が指導案の流れと合うこと（使うシートの記載）
    intro = " ".join(str(c.value) for row in ws_["はじめに"].iter_rows() for c in row if c.value)
    need = ["ステップ1_温度", S1B, S1C, "ステップ2_海と陸", "ステップ4_地域を選ぶ", "ステップ4b_スコア", "上書き保存しない", "編集を有効にする"]
    record(all(n in intro for n in need) and "ステップ3" not in intro and "ステップ5" not in intro,
           "学習者版：『はじめに』は使うシート（1・1b・1c・2・4・4b_スコア）と操作の決まりだけを案内し、使わないシートを挙げない")
    record("月全体（線の内側）" in intro and "C5" in intro and "C4" in intro and intro.index("C5") < intro.index("C4"),
           "学習者版：『はじめに』にステップ4の入力の順（C5 → C4 → 重み）と、『月全体（線の内側）』の案内が 1 行ある")
    t4a = " ".join(str(c.value) for row in ws_[S4].iter_rows(max_row=5) for c in row if c.value)
    record("ミッション" not in t4a and "ステップ4：採点する範囲を選んで" in t4a and "ミッション" not in str(ws_[S4B]["A1"].value),
           "学習者版：ステップ4・4b の題から『ミッション』を除いた")
    # 条件付き書式・チャート（Excel 保存版でも保たれていること）
    for label, wb in (("教員版", wt), ("学習者版", ws_)):
        cf4 = [str(r.sqref) for r in wb[S4].conditional_formatting]
        cf4b = [str(r.sqref) for r in wb[S4B].conditional_formatting]
        record("A6" in cf4 and "A13" in cf4b, f"{label}：状態表示の条件付き書式（ステップ4 A6、4b A13）")
        record(len(wb[S1]._charts) == 1 and len(wb["ステップ4b_南極"]._charts) == 1,
               f"{label}：チャート（ステップ1・4b_南極 に各1個）")
        record(len(wb[S1B]._charts) == 2 and len(wb[S1C]._charts) == 1,
               f"{label}：新シートのチャート（1b に2個＝日較差・夜に最高の割合、1c に1個＝24時間カーブ）")
    record("指示されたところの黄色いセルだけ" in intro and "4つの緯度帯" in intro,
           "学習者版：『はじめに』は「指示されたところの黄色いセルだけ」「ステップ1の4つの緯度帯は変えない」と書いている")
    h = " ".join(str(c.value) for row in ws_[S1]["A6:A6"] for c in row)
    record("変えずに" in h, "学習者版：ステップ1の見出し A6 は黄色い緯度帯を変えない案内（教員版は従来どおり）")
    check_structure_dai2ji(wt, ws_)
    # H 列（上位10の表の夜の最低温度）
    for label, wb in (("教員版", wt), ("学習者版", ws_)):
        w4 = wb[S4]
        record(w4["H25"].value == "夜の最低温度 [K]" and all(str(w4[f"H{r}"].value).startswith("=IFERROR(INDEX(データ_環境!$F$2:$F$14401,MATCH(")
                                                           for r in range(26, 36)),
               f"{label}：ステップ4 の上位10の表に H 列『夜の最低温度 [K]』（H26:H35＝データ_環境 F 列を引く）")
    # 文書情報に個人名がない
    with zipfile.ZipFile(student) as z:
        core = z.read("docProps/core.xml").decode("utf-8")
    with zipfile.ZipFile(teacher) as z:
        core_t = z.read("docProps/core.xml").decode("utf-8")
    ok_meta = all("lastModifiedBy" not in x and not re.search(r"<dc:creator>[^<]+</dc:creator>", x) for x in (core, core_t))
    record(ok_meta, "両版：文書情報に最終更新者・作成者がない（空）")



def check_structure_dai2ji(wt, ws_):
    """第2時の再設計（設計 §7.1）：新シート・黄色セル・入力規則・式・旧Bの残存・データ表の見出し"""
    for label, wb in (("教員版", wt), ("学習者版", ws_)):
        b = wb[S1B]
        dvs = b.data_validations.dataValidation
        lst = [d for d in dvs if d.type == "list" and "B6" in str(d.sqref) and d.formula1 == '"30,5,1"']
        whole = [d for d in dvs if d.type == "whole" and "B7" in str(d.sqref) and d.formula1 == "50" and d.formula2 == "90"]
        record(len(lst) == 1 and len(whole) == 1,
               f"{label}：1b の入力規則（B6＝帯の幅 30・5・1 の一覧、B7＝線 50〜90 の整数）")
        c = wb[S1C]
        dvc = [d for d in c.data_validations.dataValidation if d.type == "list" and "B5" in str(d.sqref) and d.formula1 == "SiteNames"]
        ref = wb.defined_names.get("SiteNames")
        record(len(dvc) == 1 and ref is not None and ref.attr_text == "データ_地点比較!$A$2:$A$6",
               f"{label}：1c の B5 のプルダウン（名前 SiteNames＝データ_地点比較!A2:A6、5択）")
        # グラフの位置：「夜に最高の割合」（縦軸 0〜8）を最初の画面（G5）に、日較差（縦軸 0〜300）を下（G15）に置く
        pos = {}
        for ch in b._charts:
            top = ch.anchor._from
            pos[ch.y_axis.scaling.max] = (top.col, top.row)       # 0 始まり：G=6、5行目=4
        record(pos.get(8) == (6, 4) and pos.get(300) == (6, 14),
               f"{label}：1b のグラフの位置（夜に最高の割合＝G5、日較差＝G15。B6・B7 を変えながら割合のグラフが最初の画面に見える）", str(pos))
        n8 = str(b["A8"].value or "")
        record("B9" in n8 and "1°ごと" in n8 and "30" in n8 and "幅を 1" in n8 and b["A8"].coordinate in b.merged_cells,
               f"{label}：1b A8 の注（B9 は1°ごとの行の値。幅30では平均で小さく見え、幅1ではB9と同じ値）に、判断・原因の語がない", n8)
        c3 = str(c["A3"].value or "")
        record("四角" in c3 and "半径の半分の長さ＋0.25度" in c3 and "半径の半分以内" not in c3,
               f"{label}：1c A3 の「内部」の説明が実装（四角＋0.25度）と一致（円の記述が残っていない）", c3)
        f9, f10 = str(b["B9"].value), str(b["B10"].value)
        record("SUMPRODUCT(MAX(" in f9 and "データ_緯度行!$B$2:$B$91<=$B$7" in f9 and "SIN(RADIANS($B$7))" in f10
               and "MAXIFS" not in f9 and "MINIFS" not in f9,
               f"{label}：1b の式（B9＝SUMPRODUCT(MAX(…))、B10＝(1−SIN(RADIANS(線)))×100。MINIFS・MAXIFS は使わない）")
        record(b["B6"].value == 30 and b["B7"].value == 90 and
               str(b["B15"].value).startswith("=AVERAGEIFS(データ_緯度行!$D$2:$D$91") and
               b.max_row >= 104,
               f"{label}：1b の既定値（幅30・線90）と、90行の表（A15:E104）の式")
        record(any("B15" in str(r.sqref) or "A15" in str(r.sqref) for r in b.conditional_formatting) and
               any("A11" in str(r.sqref) for r in b.conditional_formatting) and
               any("A6" in str(r.sqref) for r in c.conditional_formatting),
               f"{label}：1b の線の外側（灰色）と状態表示（1b A11・1c A6）の条件付き書式")
        d1 = wb[S1]["D8"].value
        record(all(str(wb[S1][f"D{r}"].value).startswith("=ROUND(AVERAGEIFS(データ_緯度行!$D$2:$D$91") for r in range(8, 12))
               and [(wb[S1][f"A{r}"].value, wb[S1][f"B{r}"].value) for r in range(8, 12)] == [(0, 6), (24, 36), (54, 66), (78, 90)],
               f"{label}：ステップ1 A8:D11 は データ_緯度行 から求める（緯度の絶対値 0〜6・24〜36・54〜66・78〜90）")
        s1 = wb[S1]
        record(s1["B17"].value == 0.25 and s1["B18"].value == 0.25 and str(s1["B22"].value).startswith("=SUMIFS(データ_温度")
               and len(s1._charts) == 1 and s1["A52"].value == "Apollo 11",
               f"{label}：旧B（B17:B18・24時間カーブ B22:B45・チャート）と旧C（着陸地点）は残っている")
        txt1 = " ".join(str(x.value) for row in s1.iter_rows() for x in row if x.value)
        record("信じてよい" not in txt1 and "極付近の値は信じ" not in txt1 and "先生のデモ" in txt1 and "帯の幅を変えたとき" in txt1,
               f"{label}：ステップ1 の A13・A47 から「信じてよい？」を除き、「ステップ1b」「先生のデモ」の案内にした")
        # 学習者版・教員版の新シートの黄色セルの既定値
    record(ws_[S1C]["B5"].value in (None, "") and wt[S1C]["B5"].value == "コペルニクス",
           "1c の B5：学習者版は空欄（担当を選ぶ）、教員版は教員の例＝コペルニクス")
    # 新しい2つのデータ表の見出し・中身
    for label, wb in (("教員版", wt), ("学習者版", ws_)):
        h1 = [str(wb["データ_緯度行"].cell(row=1, column=c).value) for c in range(1, 6)]
        h2 = [str(wb["データ_地点比較"].cell(row=1, column=c).value) for c in range(1, wb["データ_地点比較"].max_column + 1)]
        bad = [w for w in HEADER_BAD_WORDS if any(w in h for h in h1 + h2)]
        record(not bad and "21〜3時" in h1[4] and len(h2) == 58 and wb["データ_緯度行"].max_row == 91 and wb["データ_地点比較"].max_row == 6,
               f"{label}：データ_緯度行（90行×5列・見出しに窓 21〜3時）・データ_地点比較（5行×58列）の見出しに判断・原因の語がない",
               "" if not bad else f"見出しの語 {bad}")


# ---------------------------------------------------------------------------
# 3. pandas による式の再現
# ---------------------------------------------------------------------------
NORM4 = ["norm_sun_high", "norm_amp_low", "norm_earth_high", "norm_earth_low", "norm_night_warm"]
NORM4B = ["norm_illum", "norm_near_shadow", "norm_low_psf", "norm_low_slope"]
EPS = 1e-12
# 式の検査例（旧『標準の重み』。ミッション別の呼び名は廃止。重みの組は従来どおり。1位は変更前の方式と同じであることを確かめる）
STD4 = {
    "式の例1（裏側・赤道 0,2,0,3,0）": (REG_FAR_EQ, (0, 2, 0, 3, 0)),
    "式の例2（赤道の海 2,2,0,0,1）": ("赤道の海（静かの海）", (2, 2, 0, 0, 1)),
    "式の例3（赤道の海 1,2,1,0,1）": ("赤道の海（静かの海）", (1, 2, 1, 0, 1)),
}
K_STD4_FAR = "式の例1（裏側・赤道 0,2,0,3,0）"
STD4B = {"式の例1（4b 0,3,0,2）": (0, 3, 0, 2), "式の例2（4b 3,0,0,2）": (3, 0, 0, 2), "式の例3（4b 2,2,1,2）": (2, 2, 1, 2),
         # 要素ごとの分析と重み：水を1番にした班（4b）。電力＝日照率の行は1点。1位は設計 §3.1・付録Bの値
         "水1番（4b 1,3,0,0）": (1, 3, 0, 0), "水1番＋傾斜2番（4b 1,3,0,2）": (1, 3, 0, 2), "水1番・電力0（4b 0,3,0,0）": (0, 3, 0, 0)}
K_STD4B_ICE = "式の例1（4b 0,3,0,2）"
# 月全体（線の内側）：設定A〜D（重み＝太陽高度・日較差・仰角表・仰角裏・夜）。電力＝太陽高度の行は1点（設計 §3）
SET_ALL = {"A": (1, 2, 0, 0, 1), "B": (1, 0, 3, 0, 0), "C": (1, 2, 2, 0, 1), "D": (1, 1, 3, 0, 1)}
SET_ALL0 = {k + "0": (0,) + w[1:] for k, w in SET_ALL.items()}      # 試行(b)：電力を0にする
SWAP = {"A": "B", "B": "A", "C": "D", "D": "C"}                      # 試行(a)：1番を替える
LINES_ALL = (50, 70, 80, 85, 86, 88, 89, 90)
# 設計 §3.1・付録C の1位（線50〜88。Aだけ線89・90で北極側）。独立に pandas で再計算した値との突き合わせに使う
TOP_ALL_EXPECT = {"A": (-43.5, -11.5), "B": (-1.5, 0.5), "C": (-4.5, -8.5), "D": (-1.5, 0.5)}
TOP_ALL_A_POLAR = (88.5, -103.5)


class Expect:
    """pandas で表計算の式を再現する"""

    def __init__(self):
        self.env = pd.read_csv(DATA / "env_grid.csv")
        self.env["row"] = np.arange(2, len(self.env) + 2)
        self.ps = pd.read_csv(DATA / "polar_south_sites.csv")
        self.ps["row"] = np.arange(2, len(self.ps) + 2)
        self.temp = pd.read_csv(DATA / "temp_grid.csv")
        self.lr = pd.read_csv(DATA / "temp_lat_rows.csv")          # データ_緯度行 の中身（build_course_data.py の出力）
        self.sc = pd.read_csv(DATA / "site_compare.csv")           # データ_地点比較 の中身
        self.cr = pd.read_csv(DATA / "craters_labeled.csv")
        self.ref = pd.read_csv(DATA / "reference.csv")
        self.regions = list(pd.read_csv(DATA / "candidate_regions.csv")["name"])

    # --- ステップ1・2 ---
    BANDS = [(0, 6), (24, 36), (54, 66), (78, 90)]

    def bands_raw(self):
        """4バンドの日較差の平均（未丸め）。temp_lat_rows.csv の行（lat_lo>=下 かつ lat_hi<=上）の swing_K の平均"""
        return [float(self.lr.loc[(self.lr["lat_lo"] >= lo) & (self.lr["lat_hi"] <= hi), "swing_K"].mean())
                for lo, hi in self.BANDS]

    def bands(self):
        return [int(round(x)) for x in self.bands_raw()]

    def line_expect(self, line):
        """ステップ1b の B9・B10 の期待値：線の内側（行の上端が線以下）の割合の最大 [%]、使えなくなる月の面積 [%]"""
        mx = float(self.lr.loc[self.lr["lat_hi"] <= line, "night_peak_pct"].max())
        return round(mx, 1), round(float((1 - np.sin(np.radians(line))) * 100), 1)

    def band_rows(self, width):
        """ステップ1b の表（90行）の期待値：各行が入る帯（幅 width）の平均。(日較差, 夜に最高の割合)"""
        out = []
        for lo in self.lr["lat_lo"]:
            b_lo = (lo // width) * width
            m = (self.lr["lat_lo"] >= b_lo) & (self.lr["lat_hi"] <= b_lo + width)
            out.append((float(self.lr.loc[m, "swing_K"].mean()), float(self.lr.loc[m, "night_peak_pct"].mean())))
        return out

    _raw = None

    @classmethod
    def raw(cls):
        """同梱の生データ（0.5°格子）。csv の再現用（表計算の入力とは独立に pandas で作り直す）"""
        if cls._raw is None:
            d = pd.read_csv(ROOT_DATA / "diviner_global.csv.gz")
            cls._raw = d
        return cls._raw

    def density(self):
        ns = int((self.cr["区分"] == "海").sum())
        nl = int((self.cr["区分"] == "陸").sum())
        fs = float(self.ref.loc[0, "値"])
        fl = float(self.ref.loc[1, "値"])
        ds, dl = round(ns / fs), round(nl / fl)
        return ns, nl, ds, dl, round(dl / ds, 1), dl / ds

    def usgs(self):
        return float(self.ref.loc[4, "値"]), float(self.ref.loc[5, "値"])

    # --- ステップ4 ---
    def score4(self, region, w):
        e = self.env
        m = (e["region"] == region).values
        sc = np.full(len(e), np.nan)
        if sum(w) > 0:
            s = np.zeros(m.sum())
            for wi, col in zip(w, NORM4):          # Excel の式と同じ順に足す
                s = s + wi * e.loc[m, col].values
            sc[m] = s / max(1, sum(w))
        return sc

    def top10_old(self, sc, rows, k=10):
        """変更前の方式（LARGE＋MATCH(…,0)）：同じ値は最初の行を返す"""
        vals = sorted(sc[~np.isnan(sc)], reverse=True)[:k]
        out = []
        for v in vals:
            out.append(int(rows[np.where(sc == v)[0][0]]))
        return out

    def top10_new(self, sc, rows, k=10):
        """変更後：順位用列 ROUND(スコア,10)−ROW()×1E-12 で LARGE＋MATCH"""
        p = np.where(np.isnan(sc), np.nan, np.round(sc, 10) - rows * EPS)
        ok = np.where(~np.isnan(p))[0]
        order = ok[np.argsort(-p[ok], kind="stable")][:k]
        return [int(rows[i]) for i in order]

    def top4(self, region, w, which="new"):
        sc = self.score4(region, w)
        rows = self.env["row"].values
        f = self.top10_new if which == "new" else self.top10_old
        r = f(sc, rows)
        e = self.env.set_index("row")
        return r, [(float(e.loc[x, "lat"]), float(e.loc[x, "lon"])) for x in r], sc

    # --- ステップ4 月全体（線の内側）---
    def score_all(self, line, w):
        """データ_環境 O 列の式（月全体）：|緯度|≦線 のマスだけを、全球固定の正規化 J〜N の重みつき平均で採点（マス中心）"""
        e = self.env
        sc = np.full(len(e), np.nan)
        if sum(w) > 0:
            m = (e["lat"].abs() <= line).values
            s = np.zeros(int(m.sum()))
            for wi, col in zip(w, NORM4):          # Excel の式と同じ順に足す
                s = s + wi * e.loc[m, col].values
            sc[m] = s / max(1, sum(w))
        return sc

    def top_all(self, line, w, k=10):
        sc = self.score_all(line, w)
        rows = self.env["row"].values
        r = self.top10_new(sc, rows, k)
        e = self.env.set_index("row")
        return r, [(float(e.loc[x, "lat"]), float(e.loc[x, "lon"])) for x in r], sc

    # --- ステップ4b ---
    def score4b(self, w):
        p = self.ps
        sc = np.full(len(p), np.nan)
        s = np.zeros(len(p))
        for wi, col in zip(w, NORM4B):
            s = s + wi * p[col].values
        sc[:] = s / max(1, sum(w))
        return sc

    def top4b(self, w, which="new"):
        sc = self.score4b(w)
        rows = self.ps["row"].values
        if sum(w) == 0:
            return [], [], sc
        f = self.top10_new if which == "new" else self.top10_old
        r = f(sc, rows)
        e = self.ps.set_index("row")
        return r, [(float(e.loc[x, "lat"]), float(e.loc[x, "lon"])) for x in r], sc


def fmt1(x):
    """Excel の TEXT(x,"0.0") と同じ丸め（四捨五入＝0.5 は切り上げ。表示された小数どおりに丸める）"""
    from decimal import Decimal, ROUND_HALF_UP
    return str(Decimal(repr(float(x))).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))


def range_text4(pts):
    lat = [p[0] for p in pts]
    lon = [p[1] for p in pts]
    s = f"緯度 {fmt1(min(lat))} 〜 {fmt1(max(lat))} °　／　"
    if max(lon) - min(lon) > 180:
        return s + "経度は日付変更線をまたぐので範囲を出しません"
    return s + f"経度 {fmt1(min(lon))} 〜 {fmt1(max(lon))} °"


def range_text_all(pts, line):
    """月全体（ステップ4 B36）：緯度・経度の範囲＋散らばりの注意＋1位が線のいちばん外側のときの確認の一文"""
    lat = [p[0] for p in pts]
    lon = [p[1] for p in pts]
    t = f"緯度 {fmt1(min(lat))} 〜 {fmt1(max(lat))} °　／　経度 {fmt1(min(lon))} 〜 {fmt1(max(lon))} °"
    if max(lon) - min(lon) > 180:
        t += "　（経度が広く散らばっています。10行を1つずつ見ます）"
    if abs(lat[0]) >= line - 3:
        t += "　【確認】1位が線のいちばん外側のマスにあります。線を少し内側にして確かめます"
    return t


def n_cells_line(ex, line):
    """月全体の採点対象のマス数（|緯度|≦線）"""
    return int((ex.env["lat"].abs() <= line).sum())


def range_text4b(ex, rows):
    e = ex.ps.set_index("row").loc[rows]
    a, b, c = e["average_illumination_percent"], e["slope_deg"], e["km_to_shadow"]
    return (f"日照率 {fmt1(a.min())} 〜 {fmt1(a.max())} %　／　傾斜 {fmt1(b.min())} 〜 {fmt1(b.max())} °　／　"
            f"永久影まで {fmt1(c.min())} 〜 {fmt1(c.max())} km")



# 設計 §2.3 の5クレーター（夜の最低温度 [K]：内部、帯の平均、帯の9割の下端・上端、内部のセル数）。表示は小数1桁
SITE_EXPECT = {
    "ティコ": (114.2, 89.4, 87.2, 91.7, 15),
    "コペルニクス": (103.4, 95.4, 94.0, 97.3, 16),
    "ラングレヌス": (98.5, 95.1, 93.4, 97.1, 30),
    "プトレマイオス": (94.8, 95.2, 93.4, 97.3, 42),
    "アルフォンスス": (94.3, 94.7, 93.0, 96.5, 25),
}
# 設計 §2.1：緯度の絶対値の行ごとの「最高が21〜3時に来る割合」[%]（小数1桁）
ROW_EXPECT = {84: 0.0, 85: 1.0, 86: 6.8, 87: 4.2, 88: 3.7, 89: 1.5}
SITE_IDS = {"コペルニクス": "04-1-000623", "ティコ": "05-1-000975", "ラングレヌス": "07-1-000317",
            "プトレマイオス": "05-1-000082", "アルフォンスス": "05-1-000176"}


def site_from_raw(cid):
    """付録B の骨子で、生データから内部・帯を作り直す（build_course_data.py とは別に書いた再現）"""
    d = Expect.raw()
    LTc = [f"t_lt{h:02d}" for h in range(24)]
    cs = pd.read_csv(ROOT_DATA / "craters_subset.csv").set_index("crater_id")
    lat, lon, D = (float(cs.loc[cid, c]) for c in ("lat", "lon", "diam_km"))
    R = 1737.4
    dl = np.degrees(D / 2 * 0.5 / R)
    dlo = dl / np.cos(np.radians(lat))
    inner = (abs(d.lat - lat) <= dl + 0.25) & (abs(d.lon - lon) <= dlo + 0.25)
    dl2 = np.degrees(D / R) / np.cos(np.radians(lat))
    band = (abs(d.lat - lat) <= 1.5) & (abs(d.lon - lon) > dl2)
    A = d[LTc].values
    tmin = A.min(axis=1)
    return dict(n_inner=int(inner.sum()), n_band=int(band.sum()),
                inner_curve=A[inner.values].mean(axis=0), band_curve=A[band.values].mean(axis=0),
                inner_tmin=tmin[inner.values].mean(), band_tmin=tmin[band.values].mean(),
                p5=np.percentile(tmin[band.values], 5), p95=np.percentile(tmin[band.values], 95))


def check_dai2ji_pandas(ex):
    print("   [第2時の再設計：データ_緯度行・データ_地点比較・4バンド・線の式（生データから pandas で再現）]")
    d = Expect.raw()
    LTc = [f"t_lt{h:02d}" for h in range(24)]
    A = d[LTc].values
    swing = A.max(axis=1) - A.min(axis=1)
    pk = A.argmax(axis=1)
    night = ((pk >= 21) | (pk <= 3)).astype(float) * 100          # 窓＝現地時間 21〜3時
    lo = np.floor(d["lat"].abs().values).astype(int)
    g = pd.DataFrame({"lo": lo, "sw": swing, "nt": night}).groupby("lo").agg(n=("sw", "size"), sw=("sw", "mean"), nt=("nt", "mean"))
    lr = ex.lr.set_index("lat_lo")
    record(len(lr) == 90 and (lr["n_cells"] == 2880).all() and (g["n"] == 2880).all(),
           "データ_緯度行：90行、各行 2,880 セル（0.5°格子 259,200 セルを北南合わせて分ける）")
    e1 = np.abs(lr["swing_K"].values - g["sw"].values).max()
    e2 = np.abs(lr["night_peak_pct"].values - g["nt"].values).max()
    record(e1 < 6e-4 and e2 < 6e-4,
           "データ_緯度行：swing_K・night_peak_pct が、生データからの再計算と一致（差 0.0006 未満。保存は小数3桁）",
           f"最大差 {e1:.5f} / {e2:.5f}")
    rows = {k: round(float(lr.loc[k, "night_peak_pct"]), 1) for k in ROW_EXPECT}
    record(rows == ROW_EXPECT and (lr.loc[:84, "night_peak_pct"] == 0).all(),
           f"night_peak_pct（21〜3時）の設計 §2.1 の行：{rows}。0〜84°の行はすべて 0.0")
    band5 = float(lr.loc[85:89, "night_peak_pct"].mean())
    mx = {w: round(max(v[1] for v in ex.band_rows(w)), 1) for w in (30, 5, 1)}
    record(abs(band5 - 3.4) < 0.05 and mx == {30: 0.6, 5: 3.4, 1: 6.8},
           f"帯の幅を変えたときの最大：30°幅 {mx[30]}％・5°幅 {mx[5]}％・1°幅 {mx[1]}％（設計 0.6・3.4・6.8）、85〜90°帯 {band5:.2f}％")
    raw4 = ex.bands_raw()
    record(all(abs(a - b) < 1e-3 for a, b in zip(raw4, (294.508, 279.044, 233.94, 155.572))),
           f"4バンド（未丸め）{[round(x, 3) for x in raw4]}（設計 294.508・279.044・233.94・155.572）")
    direct = [float(np.mean([g.loc[k, "sw"] for k in range(lo_, hi_)])) for lo_, hi_ in ex.BANDS]
    record(all(abs(a - b) < 1e-3 for a, b in zip(direct, raw4)), "4バンドは、生データのセルから直接（行の等重み平均）求めた値とも一致")
    tcols = [f"t_lt{h:02d}" for h in range(24)]
    t0 = ex.temp[tcols].iloc[0]
    record(abs(float(ex.temp["t_swing_K"].iloc[0]) - float(t0.max() - t0.min())) < 0.11,
           "日較差の定義は教材の t_swing_K（temp_grid.csv）と同じ＝1日の最高−最低")
    # 線の式：(線, 内側の最大, 使えなくなる面積)。設計 §7.1 は「線85で 1.0／0.4」だが、式（行の上端<=線）では
    # 線85の内側に 85〜86° の行は入らず 0.0（1.0 になるのは線86）
    exp = {50: (0.0, 23.4), 70: (0.0, 6.0), 80: (0.0, 1.5), 85: (0.0, 0.4), 86: (1.0, 0.2), 87: (6.8, 0.1), 88: (6.8, 0.1), 90: (6.8, 0.0)}
    got = {k: ex.line_expect(k) for k in exp}
    record(got == exp, f"ステップ1b の線の式（pandas）：{got}")
    w = np.cos(np.radians(d["lat"].values))
    area = {k: round(float(w[np.abs(d['lat'].values) > k].sum() / w.sum() * 100), 2) for k in (70, 85)}
    record(abs(area[70] - 6.0) < 0.05 and abs(area[85] - 0.38) < 0.01,
           f"使えなくなる月の面積＝(1−sin(線))×100 は、0.5°セルの余弦重みの和でも一致：線70° {area[70]}％、線85° {area[85]}％")
    # 地点比較
    sc = ex.sc.set_index("site_name")
    bad = []
    for nm, cid in SITE_IDS.items():
        r = site_from_raw(cid)
        x = sc.loc[nm]
        ok = (r["n_inner"] == x["n_inner"] and r["n_band"] == x["n_band"]
              and abs(r["inner_tmin"] - x["inner_tmin"]) < 0.006 and abs(r["band_tmin"] - x["band_tmin"]) < 0.006
              and abs(r["p5"] - x["band_tmin_p5"]) < 0.006 and abs(r["p95"] - x["band_tmin_p95"]) < 0.006
              and np.abs(r["inner_curve"] - x[[f"inner_t_lt{h:02d}" for h in range(24)]].values.astype(float)).max() < 0.006
              and np.abs(r["band_curve"] - x[[f"band_t_lt{h:02d}" for h in range(24)]].values.astype(float)).max() < 0.006)
        if not ok:
            bad.append(nm)
    record(not bad and len(sc) == 5 and sc.shape[1] == 57,
           "データ_地点比較：5クレーターの内部・帯・カーブ・最低温度が生データからの再計算と一致（地点名＋57列＝58列）",
           "" if not bad else f"不一致 {bad}")
    bad = []
    for nm, (i, b, p5, p95, n) in SITE_EXPECT.items():
        x = sc.loc[nm]
        if not (round(x["inner_tmin"], 1) == i and round(x["band_tmin"], 1) == b and round(x["band_tmin_p5"], 1) == p5
                and round(x["band_tmin_p95"], 1) == p95 and int(x["n_inner"]) == n):
            bad.append((nm, round(x["inner_tmin"], 1), round(x["band_tmin"], 1), x["band_tmin_p5"], x["band_tmin_p95"], int(x["n_inner"])))
    record(not bad, "5クレーターの夜の最低温度が設計 §2.3 と一致（内部・帯の平均・帯の9割の下端上端・内部のセル数）", "" if not bad else str(bad))
    diffs = {nm: float(sc.loc[nm, "inner_tmin"] - sc.loc[nm, "band_tmin"]) for nm in SITE_EXPECT}
    record(round(diffs["ティコ"]) == 25 and round(diffs["コペルニクス"]) == 8 and round(diffs["ラングレヌス"], 1) == 3.4
           and round(diffs["プトレマイオス"], 1) == -0.3 and round(diffs["アルフォンスス"], 1) == -0.3,
           "内部−帯の平均：" + "、".join(f"{k}{v:+.1f}" for k, v in diffs.items()) + "（設計：＋25・＋8・＋3.4・−0.3・−0.3 K）")
    out = {nm: bool(sc.loc[nm, "inner_tmin"] > sc.loc[nm, "band_tmin_p95"] or sc.loc[nm, "inner_tmin"] < sc.loc[nm, "band_tmin_p5"])
           for nm in SITE_EXPECT}
    record([out[k] for k in SITE_EXPECT] == [True, True, True, False, False],
           f"（教員用の確認）内部が帯の9割の範囲の外：{out}（ティコ・コペルニクス・ラングレヌス＝外、プトレマイオス・アルフォンスス＝内）")


def n_places(pts, thr):
    """上位10の『場所の数』：緯度・経度がともに thr° 以内の点どうしをつなげた連結成分の数（経度は周期360°）"""
    n = len(pts)
    par = list(range(n))

    def find(i):
        while par[i] != i:
            par[i] = par[par[i]]
            i = par[i]
        return i

    for i in range(n):
        for j in range(i + 1, n):
            dlat = abs(pts[i][0] - pts[j][0])
            dlon = abs(pts[i][1] - pts[j][1])
            dlon = min(dlon, 360 - dlon)
            if dlat <= thr and dlon <= thr:
                par[find(i)] = find(j)
    return len({find(i) for i in range(n)})


def check_v3_pandas(ex):
    """要素ごとの分析と重み（docs/design_yoso_v3.md rev2）：月全体（線の内側）・4b の期待値を、別の書き方（pandas）で再現して確かめる"""
    print("   [月全体（線の内側）：設計 §3・付録B・付録C]")
    # マス中心の採点対象数（付録B）
    exp_n = {50: 8160, 60: 9600, 70: 11040, 75: 12000, 80: 12960, 85: 13440, 86: 13920, 87: 13920, 88: 13920, 89: 14400, 90: 14400}
    got_n = {k: n_cells_line(ex, k) for k in exp_n}
    record(got_n == exp_n, f"月全体の採点対象マス数（|緯度|≦線。マス中心）{got_n}")
    # 設定A〜Dの1位（線50〜90をすべて）
    bad = []
    top = {}
    for nm, w in SET_ALL.items():
        for ln in range(50, 91):
            r, pts, sc = ex.top_all(ln, list(w))
            top[(nm, ln)] = pts
            want = TOP_ALL_A_POLAR if (nm == "A" and ln >= 89) else TOP_ALL_EXPECT[nm]
            if pts[0] != want:
                bad.append((nm, ln, pts[0], want))
    record(not bad, "設定A〜Dの1位（線50〜90）：A＝（−43.5, −11.5）〔線89・90は北極側（88.5, −103.5）〕、B・D＝（−1.5, 0.5）、C＝（−4.5, −8.5）", str(bad[:3]))
    sameA = all(top[("A", ln)] == top[("A", 70)] for ln in range(50, 89))
    sameBCD = all(top[(nm, ln)] == top[(nm, 70)] for nm in "BCD" for ln in range(50, 91))
    record(sameA and sameBCD and top[("A", 89)] != top[("A", 70)],
           "上位10（順位つき）：Aは線50〜88で同一（89・90のみ別）、B・C・Dは線50〜90で同一（設計 §0-3・付録B）")
    # 北と南の同点（B・D）
    r, pts, sc = ex.top_all(70, list(SET_ALL["B"]))
    s1, s2 = round(float(sc[r[0] - 2]), 10), round(float(sc[r[1] - 2]), 10)
    record(pts[0] == (-1.5, 0.5) and pts[1] == (1.5, 0.5) and s1 == s2,
           "B 通信1番：（−1.5, 0.5）と（1.5, 0.5）が同点で、表示は南が先（行番号の小さい順）", f"スコア {s1}")
    # 上位10の場所の数（5°以内は同じ場所）とスコアの範囲
    places = {nm: n_places(top[(nm, 70)], 5) for nm in SET_ALL}
    rng = {}
    for nm, w in SET_ALL.items():
        r, pts, sc = ex.top_all(70, list(w))
        v = [round(float(sc[x - 2]), 3) for x in r]
        rng[nm] = (min(v), max(v))
    record(places == {"A": 8, "B": 1, "C": 2, "D": 1}, f"上位10の場所の数（5°以内＝同じ場所）{places}（設計：A 8・B 1・C 2・D 1）")
    p3 = {th: n_places(top[("A", 70)], th) for th in (3, 4, 6, 10, 15, 20)}
    record(p3 == {3: 8, 4: 8, 6: 8, 10: 7, 15: 6, 20: 6}, f"Aの上位10の場所の数の閾値依存 {p3}（設計：3〜6°で8、10°で7、15°・20°で6）")
    exp_rng = {"A": (0.478, 0.520), "B": (0.991, 1.0), "C": (0.632, 0.645), "D": (0.791, 0.798)}
    record(rng == exp_rng, f"上位10のスコアの範囲 {rng}（設計 §3.1：A 0.478〜0.520・B 0.991〜1.000・C 0.632〜0.645・D 0.791〜0.798）")
    # 試行 (a)(b)(c)
    ok_a = all(ex.top_all(ln, list(SET_ALL[SWAP[nm]]))[1][0] == top[(SWAP[nm], ln)][0] for nm in SET_ALL for ln in (70, 85))
    record(ok_a and top[("A", 70)][0] != top[("B", 70)][0] and top[("C", 70)][0] != top[("D", 70)][0],
           "試行(a) 1番を替える：A⇄B・C⇄D はすべて1位が動く（A→B＝緯度42°・経度12°、C→D＝緯度3°・経度9°）")
    b0 = {}
    for nm, w in SET_ALL0.items():
        for ln in range(50, 91):
            b0[(nm, ln)] = ex.top_all(ln, list(w))[1][0]
    tico, pol, w73, s82 = (-43.5, -11.5), (88.5, -103.5), (73.5, -11.5), (-82.5, 10.5)
    exp_a0 = all(b0[("A0", ln)] == (tico if ln <= 73 else w73 if ln <= 82 else s82 if ln <= 88 else pol) for ln in range(50, 91))
    exp_c0 = all(b0[("C0", ln)] == (tico if ln <= 88 else pol) for ln in range(50, 91))
    exp_bd0 = all(b0[("B0", ln)] == (-1.5, 0.5) and b0[("D0", ln)] == (-1.5, 0.5) for ln in range(50, 91))
    record(exp_a0 and exp_c0 and exp_bd0,
           "試行(b) 電力を0にする：A0＝線50〜73ティコ付近／74〜82（73.5, −11.5）／83〜88（−82.5, 10.5）／89・90北極側、"
           "C0＝線50〜88ティコ付近、B0・D0は動かない（設計 付録C）")
    c90 = {nm: ex.top_all(90, list(w))[1][0] for nm, w in SET_ALL.items()}
    record(c90["A"] == pol and all(c90[nm] == top[(nm, 70)][0] for nm in "BCD"),
           "試行(c) 線を90にする：Aだけ（88.5, −103.5）に動く（B・C・Dは動かない）")
    # 重みが規則から外れたとき（日較差だけを重く、夜0）は線で1位が9通りに動く
    exp_x = [((50, 61), (-43.5, -11.5)), ((62, 64), (61.5, 49.5)), ((65, 67), (-64.5, -100.5)), ((68, 70), (67.5, -157.5)),
             ((71, 76), (-70.5, -95.5)), ((77, 79), (-76.5, 57.5)), ((80, 82), (79.5, 93.5)), ((83, 88), (-82.5, 10.5)),
             ((89, 90), (88.5, -128.5))]
    ok_x = all(ex.top_all(ln, [1, 3, 0, 0, 0])[1][0] == pt for (lo, hi), pt in exp_x for ln in range(lo, hi + 1))
    record(ok_x, "重み(1,3,0,0,0)（日較差だけ重く、夜0）：線50〜90で1位が9通りの位置に動き、最外周に張り付く（設計 付録B）")
    # 電力の行だけ
    r, pts, sc = ex.top_all(70, [1, 0, 0, 0, 0])
    v = np.round(sc[~np.isnan(sc)], 10)
    record(int((v == v.max()).sum()) == 480 and pts[0] == (-1.5, -179.5),
           "電力の行だけ（線70）：480マスが同点で、1位は行順の（−1.5, −179.5）（設計 §2.6）")
    # 月全体の1位は、8つの箱のどれにも入らない（設計 §6）
    reg = pd.read_csv(DATA / "candidate_regions.csv")
    pts_all = [(-43.5, -11.5), (-1.5, 0.5), (1.5, 0.5), (-4.5, -8.5), (88.5, -103.5), (-85.6, 138.0)]
    inside = [(p_, rr["name"]) for p_ in pts_all for _, rr in reg.iterrows()
              if rr["lat_min"] <= p_[0] <= rr["lat_max"] and rr["lon_min"] <= p_[1] <= rr["lon_max"]]
    record(not inside, "月全体の1位（と4bの1位）は、8つの地域の箱のどれにも入らない", str(inside[:2]))
    # 4b
    print("   [4b：水を1番にした班（設計 §3.1）]")
    r, pts, sc = ex.top4b([1, 3, 0, 0])
    p_ = ex.ps.set_index("row").loc[r[0]]
    record(pts[0] == (-85.6, 138.0) and round(p_["average_illumination_percent"], 2) == 36.95 and round(p_["km_to_shadow"], 2) == 2.96
           and round(p_["slope_deg"], 2) == 9.02
           and range_text4b(ex, r).startswith("日照率 30.2 〜 44.4 %　／　傾斜 5.7 〜 19.2 °　／　永久影まで 2.2 〜 6.6 km"),
           "4b 水1番(1,3,0,0)：1位（−85.6, 138.0）、日照率36.95％・影まで2.96 km・傾斜9.02°、上位10の範囲 30.2〜44.4％／5.7〜19.2°／2.2〜6.6 km")
    r, pts, sc = ex.top4b([1, 3, 0, 2])
    p_ = ex.ps.set_index("row").loc[r[0]]
    record(pts[0] == (-84.4, 156.0) and round(p_["average_illumination_percent"], 2) == 37.95 and round(p_["km_to_shadow"], 2) == 4.82
           and round(p_["slope_deg"], 2) == 4.75
           and range_text4b(ex, r).startswith("日照率 30.1 〜 43.3 %　／　傾斜 3.6 〜 6.2 °　／　永久影まで 4.8 〜 7.9 km"),
           "4b 水1番＋傾斜2番(1,3,0,2)：1位（−84.4, 156.0）、37.95％・4.82 km・4.75°、範囲 30.1〜43.3％／3.6〜6.2°／4.8〜7.9 km（試行(d)）")
    record(ex.top4b([1, 3, 0, 1])[1][0] == (-84.4, 156.0) and ex.top4b([1, 3, 0, 3])[1][0] == (-84.4, 156.0),
           "4b 傾斜を1点でも3点でも1位は（−84.4, 156.0）で同じ（設計 §3.3）")
    r, pts, sc = ex.top4b([0, 3, 0, 0])
    v = np.round(sc, 10)
    record(pts[0] == (-88.4, -147.0) and int((v == v.max()).sum()) == 127 and ex.top4b([0, 3, 0, 2])[1][0] == (-88.2, 114.0),
           "4b 電力0(0,3,0,0)：127地点が同点で1位は行順の（−88.4, −147.0）。(0,3,0,2) は（−88.2, 114.0）（試行にしない。設計 §2.6）")


def check_pandas(ex):
    print("\n== 3a. pandas による式の再現 ==")
    b = ex.bands()
    record(b == [295, 279, 234, 156], f"ステップ1の4バンドの日較差の平均 {b}（期待 295・279・234・156 K。データ_緯度行 の行から）")
    check_dai2ji_pandas(ex)
    check_v3_pandas(ex)
    ns, nl, ds, dl, r1, r = ex.density()
    record((ns, nl) == (2232, 34145) and abs(r - 3.52) < 0.005,
           f"ステップ2：クレーター数 海{ns}・陸{nl}、密度 海{ds}・陸{dl}、倍率 {r:.3f}（表示は小数1桁で {r1}）")
    us, ul = ex.usgs()
    record(round(us, 2) == 3.16 and round(ul, 2) == 2.38, f"USGS の相対年代 海{us}→{round(us, 2)}・陸{ul}→{round(ul, 2)}")

    # --- 同点重複：全地域×単一指標（ステップ4） ---
    print("   [全地域×単一指標（重み 3）の上位10：異なる行数（変更前→変更後）]")
    bad_new, bad_old, same1 = [], 0, True
    names4 = ["太陽高度", "日較差", "仰角表", "仰角裏", "夜の最低"]
    ncases = 0
    for reg in ex.regions:
        n_rows = int((ex.env["region"] == reg).sum())
        for i in range(5):
            w = [0] * 5
            w[i] = 3
            sc = ex.score4(reg, w)
            rows = ex.env["row"].values
            old = ex.top10_old(sc, rows)
            new = ex.top10_new(sc, rows)
            want = min(10, n_rows)
            ncases += 1
            if len(set(old)) < want:
                bad_old += 1
            if len(set(new)) != want:
                bad_new.append((reg, names4[i], len(set(new))))
            if old[0] != new[0]:
                same1 = False
    record(not bad_new, f"ステップ4：{ncases} 通り（8地域×5指標）すべてで上位10が全部異なる行（変更後）",
           f"変更前に重複した通り数 {bad_old}／{ncases}" + (f"　不合格 {bad_new}" if bad_new else ""))
    record(same1, "単一指標の全 40 通りで、1位は変更前と同じ行")
    # 4b 単一指標
    bad_b, bad_ob = [], 0
    names4b = ["日照", "永久影まで", "永久影率", "傾斜"]
    for i in range(4):
        w = [0] * 4
        w[i] = 3
        sc = ex.score4b(w)
        rows = ex.ps["row"].values
        old = ex.top10_old(sc, rows)
        new = ex.top10_new(sc, rows)
        if len(set(old)) < 10:
            bad_ob += 1
        if len(set(new)) != 10:
            bad_b.append((names4b[i], len(set(new))))
        if old[0] != new[0]:
            bad_b.append(("1位が変わった", names4b[i]))
    record(not bad_b, "4b：単一指標 4 通りで上位10が全部異なる行（変更後）", f"変更前に重複した通り数 {bad_ob}／4")
    # 赤道の海（3,0,0,0,0）と 4b（0,3,0,0）（0,0,3,0）
    for lab, f in (("赤道の海 (3,0,0,0,0)", lambda: ex.top4("赤道の海（静かの海）", [3, 0, 0, 0, 0], "old")),):
        r, pts, sc = f()
        record(len(set(r)) < 10, f"{lab}：変更前の方式では上位10が {len(set(r))} 行しかない（不具合の再現）")
    for lab, w in (("4b (0,3,0,0)", [0, 3, 0, 0]), ("4b (0,0,3,0)", [0, 0, 3, 0])):
        r, pts, sc = ex.top4b(w, "old")
        record(len(set(r)) < 10, f"{lab}：変更前の方式では上位10が {len(set(r))} 行しかない（不具合の再現）")

    # --- 式の検査例（旧・標準の重み）---
    print("   [式の検査例（重みの組は従来どおり。1位は変更前の方式と同じ）]")
    for lab, (reg, w) in STD4.items():
        ro, po, _ = ex.top4(reg, list(w), "old")
        rn, pn, _ = ex.top4(reg, list(w), "new")
        record(ro[0] == rn[0], f"{lab}：1位は変更前と同じ地点 {pn[0]}", f"変更前の1位 {po[0]}")
    for lab, w in STD4B.items():
        ro, po, _ = ex.top4b(list(w), "old")
        rn, pn, _ = ex.top4b(list(w), "new")
        record(ro[0] == rn[0], f"{lab}：1位は変更前と同じ地点 {pn[0]}", f"変更前の1位 {po[0]}")
    # --- A4 の期待範囲 ---
    reg, w = STD4[K_STD4_FAR]
    rn, pn, _ = ex.top4(reg, list(w))
    lat = [p[0] for p in pn]
    lon = [p[1] for p in pn]
    record(max(abs(x) for x in lat) <= 4.5 and min(lon) >= 175.5 and max(lon) <= 178.5,
           f"A4 裏側・赤道（0,2,0,3,0）：上位10は緯度±4.5°以内・経度175.5〜178.5°（実測 緯度{min(lat)}〜{max(lat)}、経度{min(lon)}〜{max(lon)}）",
           range_text4(pn))
    rn, pn, _ = ex.top4b(list(STD4B[K_STD4B_ICE]))
    e = ex.ps.set_index("row").loc[rn]
    record((e["average_illumination_percent"] == 0).all() and (e["km_to_shadow"] == 0).all(),
           "A4 4b（0,3,0,2）：上位10は日照 0%・永久影までの距離 0 km", range_text4b(ex, rn))


# ---------------------------------------------------------------------------
# 3b. Excel 実機（COM）
# ---------------------------------------------------------------------------
def run_excel(job_steps, src, workdir, saveas=None):
    work = pathlib.Path(workdir)
    work.mkdir(parents=True, exist_ok=True)
    copy = work / ("copy_" + pathlib.Path(src).name)
    shutil.copyfile(src, copy)               # 元ファイルは開かない
    job = {"file": str(copy), "out": str(work / "out.json"), "steps": job_steps}
    jp = work / "job.json"
    jp.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
    r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(DRIVER),
                        "-JobFile", str(jp)], capture_output=True, text=True, timeout=1500)
    if r.returncode != 0:
        raise RuntimeError("Excel ドライバが失敗: " + (r.stderr or r.stdout)[:1000])
    return json.loads((work / "out.json").read_text(encoding="utf-8"))


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def top_from(grid):
    """C26:D35 の2列 → [(緯度, 経度)…]（空白は除く）"""
    return [(num(r[0]), num(r[1])) for r in grid if num(r[0]) is not None]



SITES = ["コペルニクス", "ティコ", "ラングレヌス", "プトレマイオス", "アルフォンスス"]
LINES_CHECK = (50, 70, 80, 85, 86, 87, 88, 90)


def _f(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def check_excel_dai2ji(ex, out, label, student, lines):
    """第2時の再設計：Excel 実機の再計算結果を pandas の再現と突き合わせる"""
    # 4バンド・旧B
    bands = [int(r[3]) for r in out["bands_after"]]
    record(bands == [295, 279, 234, 156], f"{label}：ステップ1 A8:D11（Excel）の4バンド {bands}")
    cnt = [int(r[2]) for r in out["bands_after"]]
    exp_cnt = [int(ex.lr.loc[(ex.lr["lat_lo"] >= lo) & (ex.lr["lat_hi"] <= hi), "n_cells"].sum()) for lo, hi in ex.BANDS]
    record(cnt == exp_cnt, f"{label}：ステップ1 C8:C11（セル数 0.5°）{cnt}")
    pol = [r[0] for r in out["old_b_polar"]]
    eq = [r[0] for r in out["old_b_eq"]]
    record(abs(pol[0] - 49.6) < 0.06 and abs(pol[1] - 174.4) < 0.06 and abs(pol[2] - 170.5) < 0.06 and abs(max(eq) - 394.0) < 6 and abs(min(eq) - 96.2) < 0.06,
           f"{label}：旧デモ（B17・B18＝-86.75・0.25）の24時間カーブは 0時{pol[0]}・1時{pol[1]}・2時{pol[2]} K（設計：49.6・174.4・170.5）、赤道 0.25・0.25 は最高{max(eq)}・最低{min(eq)} K（設計 約394・96.2）")
    # 入力規則（Excel 実機）
    d6, d7, d5 = out["dv_b6"], out["dv_b7"], out["dv_site"]
    record(d6.get("type") == 3 and "30,5,1" in str(d6.get("formula1")) and d6.get("in_cell_dropdown") is True and d6.get("alert_style") == 1
           and d7.get("type") == 1 and d7.get("formula1") == "50" and d7.get("formula2") == "90" and d7.get("operator") == 1
           and d5.get("type") == 3 and d5.get("formula1") == "=SiteNames" and d5.get("in_cell_dropdown") is True and d5.get("alert_style") == 1,
           f"{label}：入力規則（Excel）：1b B6＝一覧 {d6.get('formula1')}、B7＝整数 {d7.get('formula1')}〜{d7.get('formula2')}、1c B5＝一覧 {d5.get('formula1')}（いずれも停止メッセージ）")
    okv6 = [out[f"v6_ok{n}"] for n in (30, 5, 1)]
    ngv6 = [out[f"v6_ng{n}"] for n in (10, 0, "abc")]
    record(all(v is True for v in okv6) and all(v is False for v in ngv6),
           f"{label}：1b B6（帯の幅）の入力規則：30・5・1 は受理、10・0・abc は拒否")
    okv7 = [out[f"v7_ok{n}"] for n in (50, 90, 85)]
    ngv7 = [out[f"v7_ng{n}"] for n in (49, 91, "85.5", "x")]
    record(all(v is True for v in okv7) and all(v is False for v in ngv7),
           f"{label}：1b B7（線）の入力規則：50・90・85 は受理、49・91・85.5・x は拒否")
    record(all(out[f"v5_ok{i}"] is True for i in range(5)) and out["v5_ng"] is False and out["v5_ng2"] is False,
           f"{label}：1c B5（クレーター）の入力規則：5 つの名前は受理、「でたらめ」「コペ」（名前の一部）は拒否")
    # 1b：幅ごとの表（90行）と、線の式
    for wd in (30, 5, 1):
        tbl = out[f"b1_tbl_{wd}"]
        exp = ex.band_rows(wd)
        got = [(_f(r[1]), _f(r[2])) for r in tbl]
        ok = (len(got) == 90 and all(g[0] is not None and abs(g[0] - e[0]) < 1e-6 and abs(g[1] - e[1]) < 1e-6 for g, e in zip(got, exp))
              and [int(r[0]) for r in tbl] == list(range(90)))
        record(ok, f"{label}：1b の表（90行）が、帯の幅 {wd}° の pandas の再現と一致（日較差・最高が21〜3時に来る割合）",
               f"最大の割合 {max(g[1] for g in got):.2f}％")
        res = [r[0] for r in out[f"b1_res_{wd}"]]
        st = out[f"b1_st_{wd}"][0][0]
        record(res[0] == 6.8 and res[1] == 0.0 and st.startswith("OK"),
               f"{label}：幅{wd}・線90 の B9・B10＝{res}、状態表示「{st}」")
    bad = []
    for ln in lines:
        r = [x[0] for x in out[f"b1_line_{ln}"]]
        e = ex.line_expect(ln)
        if not (abs(r[0] - e[0]) < 1e-9 and abs(r[1] - e[1]) < 1e-9):
            bad.append((ln, r, e))
    record(not bad, f"{label}：1b の B9・B10（線 {list(lines)}）が pandas の再現と一致（線70＝0.0／6.0、線85＝0.0／0.4、線86＝1.0／0.2、線87＝6.8／0.1）",
           "" if not bad else str(bad))
    b9 = out["b1_bad7_res"]
    record(all(x[0] == "" for x in b9) and out["b1_bad7_st"][0][0].startswith("【注意】線は 50〜90 の整数"),
           f"{label}：1b 線が範囲外（貼り付けで 95）→ B9・B10 は空白、「{out['b1_bad7_st'][0][0]}」")
    record(out["b1_bad6_st"][0][0].startswith("【注意】帯の幅は 30・5・1"),
           f"{label}：1b 帯の幅が範囲外（貼り付けで 10）→「{out['b1_bad6_st'][0][0]}」")
    # 1c：5クレーター
    sc = ex.sc.set_index("site_name")
    bad, bad_c = [], []
    for i, nm in enumerate(SITES):
        row = [_f(x) for x in out[f"c1_row_{i}"][0]]
        x = sc.loc[nm]
        exp = [round(x["inner_tmin"], 1), round(x["band_tmin"], 1), round(x["band_tmin_p5"], 1), round(x["band_tmin_p95"], 1)]
        if row != exp:
            bad.append((nm, row, exp))
        curve = [(_f(a[0]), _f(a[1])) for a in out[f"c1_curve_{i}"]]
        ei = [float(x[f"inner_t_lt{h:02d}"]) for h in range(24)]
        eb = [float(x[f"band_t_lt{h:02d}"]) for h in range(24)]
        if not all(abs(c[0] - a) < 1e-6 and abs(c[1] - b) < 1e-6 for c, a, b in zip(curve, ei, eb)):
            bad_c.append(nm)
        st = out[f"c1_st_{i}"][0][0]
        if not st.startswith(f"OK　{nm}：内部 {int(x['n_inner'])} セル、帯 {int(x['n_band'])} セル"):
            bad.append((nm, st))
    record(not bad, f"{label}：1c の夜の最低温度の1行（内部・帯の平均・帯の9割の下端上端）と状態表示が、5クレーターとも pandas の再現と一致", "" if not bad else str(bad))
    record(not bad_c, f"{label}：1c の24時間カーブ（B15:C38）が、5クレーターとも データ_地点比較 と一致", "" if not bad_c else str(bad_c))
    record(out["c1_blank_st"][0][0].startswith("担当のクレーターを選んでください") and all(x == "" for x in out["c1_blank_row"][0])
           and out["c1_bad_st"][0][0].startswith("【注意】B5 の名前が") and all(x == "" for x in out["c1_bad_row"][0]),
           f"{label}：1c B5 が空欄 →「{out['c1_blank_st'][0][0]}」、不一致 →「{out['c1_bad_st'][0][0]}」（表は空白）")
    # チャート（Excel が描く）
    sheets = {n: c for n, v, c in out["info"]["sheets"]}
    record(sheets.get(S1) == 1 and sheets.get(S1B) == 2 and sheets.get(S1C) == 1,
           f"{label}：Excel のチャート数：ステップ1＝{sheets.get(S1)}（旧B のまま）、1b＝{sheets.get(S1B)}、1c＝{sheets.get(S1C)}")
    c1b, c1c = out["chart_1b"], out["chart_1c"]
    ax1 = [c[5] for c in c1b] if c1b and len(c1b[0]) > 5 else []
    ok1b = (len(c1b) == 2 and len(ax1) == 2 and ax1[0].get("y_min") == 0 and ax1[0].get("y_max") == 300 and ax1[1].get("y_max") == 8
            and c1b[1][1] == "$G$5" and c1b[0][1] == "$G$15"
            and all(a.get("x_label_spacing") == 10 and a.get("legend") is False for a in ax1)
            and all(len(c[4]) == 1 for c in c1b))
    ax2 = c1c[0][5] if c1c and len(c1c[0]) > 5 else {}
    ok1c = len(c1c) == 1 and len(c1c[0][4]) == 2 and ax2.get("legend") is True
    record(ok1b and ok1c,
           f"{label}：チャートの軸・系列（Excel）1b＝縦軸 {ax1[0].get('y_min') if ax1 else '?'}〜{ax1[0].get('y_max') if ax1 else '?'} と 0〜{ax1[1].get('y_max') if ax1 else '?'}（固定）・横軸の目盛 10 ごと、1c＝2系列・凡例あり",
           f"1b {c1b[0][1] if c1b else ''}・{c1b[1][1] if len(c1b) > 1 else ''}、1c {c1c[0][1] if c1c else ''}")
    record((out.get("png_1b") == 2 and out.get("png_1c") == 1) and not out.get("errors"),
           f"{label}：Excel が描いたチャートを PNG に書き出せた（1b 2個・1c 1個。目視確認用）")

def check_excel_v3(ex, out, label, student):
    """要素ごとの分析と重み（設計 §2.5・§2.9、AC9・AC10・AC11）：ステップ4『月全体（線の内側）』の Excel 実機検査"""
    # --- プルダウンの一覧（非表示列 M）---
    mcol = [r[0] for r in out["m_col"]]
    names = {n: ref for n, ref in out["info"]["names"]}
    record(out["m_hidden"] is True and mcol[0] == ALL_MOON and mcol[1:9] == ex.regions and mcol[9] == ""
           and names.get("AreaChoices", "").endswith("$M$5:$M$13") and names.get("AllMoon", "").endswith("$M$5")
           and names.get("LineOK", "").endswith("$M$14"),
           f"{label}：Excel：非表示列 M の一覧は先頭が『{ALL_MOON}』＋8地域、名前 AreaChoices・AllMoon・LineOK が M5:M13・M5・M14 を指す",
           f"M14（線が未入力のとき）＝『{mcol[9]}』")
    d4 = out["dv_c4"]
    okv = [out[f"v4_ok{n}"] for n in (50, 85, 90)]
    ngv = [out[f"v4_ng{n}"] for n in ("49", "91", "85.5", "x")]
    record(d4.get("type") == 1 and d4.get("formula1") == "50" and d4.get("formula2") == "90" and d4.get("operator") == 1
           and d4.get("alert_style") == 1 and all(v is True for v in okv) and all(v is False for v in ngv),
           f"{label}：Excel：C4（線）の入力規則＝整数 50〜90（停止メッセージ）。50・85・90 は受理、49・91・85.5・x は拒否")
    # --- 月全体：設定A〜D・A0〜D0（電力0）× 線50・70・80・85・86・88・89・90 ---
    bad_top, bad_rng, bad_st, bad_h, n_ok, fired = [], [], [], [], 0, {}
    for nm, w in list(SET_ALL.items()) + list(SET_ALL0.items()):
        for ln in LINES_ALL:
            key = f"all_{nm}_{ln}"
            r, pts, sc = ex.top_all(ln, list(w))
            got = top_from(out[key + "_top"])
            if got != pts:
                bad_top.append((key, got[:2], pts[:2]))
            rng = range_text_all(pts, ln)
            if out[key + "_rng"][0][0] != rng:
                bad_rng.append((key, out[key + "_rng"][0][0], rng))
            fired[(nm, ln)] = "【確認】" in out[key + "_rng"][0][0]
            st = (f"OK　月全体（線 {ln}° の内側）　採点対象 {n_cells_line(ex, ln)} 行　（重みの合計 {sum(w)}）"
                  + ("　※太陽高度の行が0です" if w[0] == 0 else ""))
            if out[key + "_st"][0][0] != st:
                bad_st.append((key, out[key + "_st"][0][0], st))
            hh = [_f(x[0]) for x in out[key + "_h"]]
            e = ex.env.set_index("row")
            exp_h = [float(e.loc[x, "night_min_K"]) for x in r]
            if not (len(hh) == 10 and all(a is not None and abs(a - b) < 1e-9 for a, b in zip(hh, exp_h))):
                bad_h.append(key)
            n_ok += 1
    record(not bad_top, f"{label}：Excel：月全体の上位10の緯度・経度が pandas の再現と一致（設定A〜D・A0〜D0 × 線 {list(LINES_ALL)} の {n_ok} 通り）",
           "" if not bad_top else str(bad_top[:3]))
    record(not bad_st, f"{label}：Excel：月全体の状態表示『OK　月全体（線 □° の内側）　採点対象 N 行　（重みの合計 □）』が全 {n_ok} 通りで一致"
           "（線70＝11040行・85＝13440行・90＝14400行。電力0では『※太陽高度の行が0です』の付記）", "" if not bad_st else str(bad_st[:2]))
    record(not bad_rng, f"{label}：Excel：月全体の範囲表示 B36（緯度・経度の範囲、散らばりの注意、1位が線の最外周のときの【確認】）が全 {n_ok} 通りで一致",
           "" if not bad_rng else str(bad_rng[:2]))
    record(not bad_h, f"{label}：Excel：上位10の H 列『夜の最低温度 [K]』が データ_環境 F 列の値と一致（全 {n_ok} 通り）")
    # --- B36 の確認の一文が出る／出ない ---
    main_quiet = all(not fired[(nm, ln)] for nm in SET_ALL for ln in LINES_ALL if not (nm == "A" and ln >= 89))
    a_polar = fired[("A", 89)] and fired[("A", 90)]
    x_fired = all("【確認】" in out[f"all_X_{ln}_rng"][0][0] for ln in (80, 85))
    x_first = [top_from(out[f"all_X_{ln}_top"])[0] for ln in (80, 85)]
    record(main_quiet and a_polar and x_fired and x_first == [(79.5, 93.5), (-82.5, 10.5)],
           f"{label}：Excel：B36 の確認の一文は、本線（A〜D・線50〜88）では出ず、Aの線89・90と重み(1,3,0,0,0)の線80・85（1位が最外周 79.5°・82.5°）で出る")
    # --- 設計の1位（pandas の再現とは別に、設計の数字そのもの）---
    for nm, w in SET_ALL.items():
        got70 = top_from(out[f"all_{nm}_70_top"])[0]
        got85 = top_from(out[f"all_{nm}_85_top"])[0]
        if not (got70 == TOP_ALL_EXPECT[nm] and got85 == TOP_ALL_EXPECT[nm]):
            record(False, f"{label}：Excel：設定 {nm} の1位（線70・85）が設計の値と違う", f"{got70}/{got85}")
            break
    else:
        record(top_from(out["all_A_89_top"])[0] == TOP_ALL_A_POLAR and top_from(out["all_A_90_top"])[0] == TOP_ALL_A_POLAR,
               f"{label}：Excel：1位 A＝（−43.5, −11.5）、B・D＝（−1.5, 0.5）、C＝（−4.5, −8.5）〔線70・85〕、A は線89・90で北極側（88.5, −103.5）")
    # --- 試行 (a)(b)(c)（Excel の1位）---
    t = lambda nm, ln: top_from(out[f"all_{nm}_{ln}_top"])[0]
    ok_a = all(t(SWAP[nm], ln) != t(nm, ln) for nm in SET_ALL for ln in (70, 85))
    ok_b = (t("A0", 70) == t("A", 70) and t("A0", 80) == (73.5, -11.5) and t("A0", 85) == (-82.5, 10.5)
            and t("C0", 70) == (-43.5, -11.5) and t("B0", 70) == t("B", 70) and t("D0", 70) == t("D", 70))
    ok_c = t("A", 90) == TOP_ALL_A_POLAR and all(t(nm, 90) == t(nm, 70) for nm in "BCD")
    record(ok_a and ok_b and ok_c,
           f"{label}：Excel：試行(a) 1番を替える＝全設定で動く／(b) 電力を0にする＝A0は線70で動かず・80で（73.5, −11.5）・85で（−82.5, 10.5）、C0は（−43.5, −11.5）、B0・D0は動かない／"
           "(c) 線を90にする＝Aだけ北極側へ")
    # --- 負の検査（AC11）---
    n = lambda k: out["neg_" + k + "_st"][0][0]
    top_blank = lambda k: all(r[0] == "" for r in out["neg_" + k + "_top"])
    for k in ("c4_blank", "c4_95", "c4_49", "c4_dec", "c4_text"):
        record(n(k).startswith("【注意】月全体で採点するには、C4 に線（50〜90 の整数。第2時で引いた線）を入れてください")
               and top_blank(k) and out["neg_" + k + "_rng"][0][0] == "",
               f"{label}：Excel：C4 が空欄・範囲外・小数・文字のとき（{k}）→ 状態表示が止め、上位10と範囲表示は空白 『{n(k)[:30]}…』")
    record(n("only_power").startswith("【注意】太陽高度の行だけでは、同じ点数の場所が大量に出ます"),
           f"{label}：Excel：電力（C8）の行だけに重み → 『{n('only_power')}』（月全体のみの注意。1位は行順の（−1.5, −179.5））")
    ot = top_from(out["neg_only_power_top"])
    record(ot and ot[0] == (-1.5, -179.5), f"{label}：Excel：電力の行だけのときの1位は行順の（−1.5, −179.5）（意味のない1位。注意で止める）")
    record(n("zero").startswith("【注意】重みがすべて0です") and top_blank("zero"),
           f"{label}：Excel：重みがすべて0 → 『{n('zero')}』、上位10は空白")
    record(n("c8_0").startswith("OK　月全体（線 70° の内側）") and "【注意】" not in n("c8_0") and n("c8_0").endswith("※太陽高度の行が0です"),
           f"{label}：Excel：電力（C8）＝0 → 注意色の【注意】でなく OK の付記『{n('c8_0')[-12:]}』")
    record(n("over").startswith("【注意】重みは0〜5の整数で入れてください"), f"{label}：Excel：重みが範囲外（貼り付けで 9）→ 『{n('over')}』")
    # --- 地域モードに戻したとき不変 ---
    rg = top_from(out["back_region_top"])
    r, pts, sc = ex.top4("赤道の海（静かの海）", [1, 2, 0, 0, 1])
    record(rg == pts and rg[0] == (1.5, 33.5) and out["back_region_st"][0][0].startswith("OK　採点対象 65 行")
           and out["back_region_rng"][0][0] == range_text4(pts),
           f"{label}：Excel：月全体のあとに地域モード（赤道の海・設定A）へ戻すと、1位（1.5, 33.5）・状態表示・範囲表示が従来どおり")


def check_excel(ex, path, label, workdir, student):
    print(f"\n== 3b. Excel 実機（{label}） ==")
    steps = [{"op": "info", "key": "info"}, {"op": "calc", "key": "first"}]
    g = lambda sh, addr, key, text=False: steps.append({"op": "get", "sheet": sh, "addr": addr, "key": key, "text": text})
    g(S1, "A8:D11", "bands")
    g(S2, "A7:B13", "s2")
    g(S2, "A26:B28", "s2d")
    g("ステップ4_地域を選ぶ", "A6", "st_init", True)
    steps.append({"op": "dv", "sheet": S4, "addr": "C5", "key": "dv_c5"})
    steps.append({"op": "dv", "sheet": S4, "addr": "C8", "key": "dv_w4"})
    steps.append({"op": "dv", "sheet": S4B, "addr": "C8", "key": "dv_w4b"})
    # 入力規則の動作（拒否されるか）
    tests_w = [("w_ok3", 3), ("w_ok0", 0), ("w_ok5", 5), ("w_ng-1", -1), ("w_ng9", 9), ("w_ng2.5", 2.5), ("w_ngabc", "abc")]
    for k, v in tests_w:
        steps.append({"op": "testvalid", "sheet": S4, "addr": "C8", "value": v, "key": k})
        steps.append({"op": "testvalid", "sheet": S4B, "addr": "C8", "value": v, "key": k + "_b"})
    for i, nm in enumerate(ex.regions):
        steps.append({"op": "testvalid", "sheet": S4, "addr": "C5", "value": nm, "key": f"c5_ok{i}"})
    steps.append({"op": "testvalid", "sheet": S4, "addr": "C5", "value": "でたらめ", "key": "c5_ng"})
    steps.append({"op": "testvalid", "sheet": S4, "addr": "C5", "value": "赤道の海", "key": "c5_ng2"})
    steps.append({"op": "testvalid", "sheet": S4, "addr": "C5", "value": ALL_MOON, "key": "c5_all_ok"})
    steps.append({"op": "testvalid", "sheet": S4, "addr": "C5", "value": "月全体", "key": "c5_all_ng"})
    steps.append({"op": "dv", "sheet": S4, "addr": "C4", "key": "dv_c4"})
    for k, v in (("ok50", 50), ("ok85", 85), ("ok90", 90), ("ng49", 49), ("ng91", 91), ("ng85.5", 85.5), ("ngx", "x")):
        steps.append({"op": "testvalid", "sheet": S4, "addr": "C4", "value": v, "key": "v4_" + k})
    steps.append({"op": "colhidden", "sheet": S4, "col": "M", "key": "m_hidden"})
    g(S4, "M5:M14", "m_col", True)

    cases = []   # (key, 種別, 地域, 重み)
    # ステップ4：全地域×単一指標（重み3）
    for ri, reg in enumerate(ex.regions):
        for i in range(5):
            w = [0] * 5
            w[i] = 3
            cases.append((f"s4_{ri}_{i}", 4, reg, w))
    # 式の検査例と、重みが全部0／1の例
    for lab, (reg, w) in STD4.items():
        cases.append((f"std4_{lab}", 4, reg, list(w)))
    cases.append(("s4_all1", 4, REG_FAR_EQ, [1, 1, 1, 1, 1]))
    cases.append(("s4_zero", 4, "赤道の海（静かの海）", [0, 0, 0, 0, 0]))
    # 4b
    for i in range(4):
        w = [0] * 4
        w[i] = 3
        cases.append((f"s4b_single{i}", "4b", None, w))
    for lab, w in STD4B.items():
        cases.append((f"std4b_{lab}", "4b", None, list(w)))
    cases.append(("s4b_zero", "4b", None, [0, 0, 0, 0]))
    cases.append(("s4b_over", "4b", None, [3, 0, 0, 9]))

    for key, kind, reg, w in cases:
        if kind == 4:
            steps.append({"op": "set", "sheet": S4, "addr": "C5", "value": reg})
            for j, v in enumerate(w):
                steps.append({"op": "set", "sheet": S4, "addr": f"C{8 + j}", "value": v})
            steps.append({"op": "calc", "key": key})
            g(S4, "C26:D35", key + "_top")
            g(S4, "B26:B35", key + "_score", True)
            g(S4, "A6", key + "_st", True)
            g(S4, "B36", key + "_rng", True)
        else:
            for j, v in enumerate(w):
                steps.append({"op": "set", "sheet": S4B, "addr": f"C{8 + j}", "value": v})
            steps.append({"op": "calc", "key": key})
            g(S4B, "C26:D35", key + "_top")
            g(S4B, "A13", key + "_st", True)
            g(S4B, "B36", key + "_rng", True)
    # 各地域を C5 に入れたときの状態表示（重みは標準）
    for ri, reg in enumerate(ex.regions):
        steps.append({"op": "set", "sheet": S4, "addr": "C5", "value": reg})
        for j, v in enumerate([2, 2, 0, 0, 0]):
            steps.append({"op": "set", "sheet": S4, "addr": f"C{8 + j}", "value": v})
        steps.append({"op": "calc", "key": f"reg{ri}"})
        g(S4, "A6", f"reg{ri}_st", True)
        g(S4, "C26:C35", f"reg{ri}_lat")
    for key, val in (("c5_blank", None), ("c5_bad", "でたらめ")):
        steps.append({"op": "set", "sheet": S4, "addr": "C5", "value": val})
        steps.append({"op": "calc", "key": key})
        g(S4, "A6", key + "_st", True)
        g(S4, "C26:C35", key + "_lat", True)
    # ---- 要素ごとの分析と重み：ステップ4『月全体（線の内側）』（設計 §2.5・§3、受け入れ基準 AC9・AC11）----
    steps.append({"op": "set", "sheet": S4, "addr": "C5", "value": ALL_MOON})
    all_cases = []
    for nm, w in list(SET_ALL.items()) + list(SET_ALL0.items()):
        for ln in LINES_ALL:
            all_cases.append((f"all_{nm}_{ln}", ln, w))
    for ln in (80, 85):       # 規則から外れた重み（日較差だけ重く、夜0）：B36 の確認の一文が出る
        all_cases.append((f"all_X_{ln}", ln, (1, 3, 0, 0, 0)))
    for key, ln, w in all_cases:
        steps.append({"op": "set", "sheet": S4, "addr": "C4", "value": ln})
        for j, v in enumerate(w):
            steps.append({"op": "set", "sheet": S4, "addr": f"C{8 + j}", "value": v})
        steps.append({"op": "calc", "key": key})
        g(S4, "C26:D35", key + "_top")
        g(S4, "H26:H35", key + "_h")
        g(S4, "A6", key + "_st", True)
        g(S4, "B36", key + "_rng", True)
    # 負の検査：C4 が空欄・範囲外（貼り付けで 95）・文字、電力の行だけ、重みが全部0、C8＝0（付記）、重みが範囲外
    neg_cases = [("c4_blank", None, (1, 2, 0, 0, 1)), ("c4_95", 95, (1, 2, 0, 0, 1)), ("c4_49", 49, (1, 2, 0, 0, 1)),
                 ("c4_dec", 85.5, (1, 2, 0, 0, 1)), ("c4_text", "abc", (1, 2, 0, 0, 1)),
                 ("only_power", 70, (1, 0, 0, 0, 0)), ("zero", 70, (0, 0, 0, 0, 0)), ("c8_0", 70, (0, 2, 0, 0, 1)),
                 ("over", 70, (1, 9, 0, 0, 1))]
    for key, c4, w in neg_cases:
        steps.append({"op": "set", "sheet": S4, "addr": "C4", "value": c4})
        for j, v in enumerate(w):
            steps.append({"op": "set", "sheet": S4, "addr": f"C{8 + j}", "value": v})
        steps.append({"op": "calc", "key": "neg_" + key})
        g(S4, "A6", "neg_" + key + "_st", True)
        g(S4, "C26:D35", "neg_" + key + "_top", True)
        g(S4, "B36", "neg_" + key + "_rng", True)
    # 地域モードへ戻して不変を確かめる（月全体のあとに赤道の海・設定A。以前のケースと別の順序で）
    steps.append({"op": "set", "sheet": S4, "addr": "C4", "value": 70})
    steps.append({"op": "set", "sheet": S4, "addr": "C5", "value": "赤道の海（静かの海）"})
    for j, v in enumerate((1, 2, 0, 0, 1)):
        steps.append({"op": "set", "sheet": S4, "addr": f"C{8 + j}", "value": v})
    steps.append({"op": "calc", "key": "back_region"})
    g(S4, "C26:D35", "back_region_top")
    g(S4, "A6", "back_region_st", True)
    g(S4, "B36", "back_region_rng", True)
    steps.append({"op": "chart", "sheet": S1, "key": "chart1"})
    steps.append({"op": "chart", "sheet": "ステップ4b_南極", "key": "chart2"})

    # ---- 第2時の再設計：ステップ1（旧B）・1b・1c ----
    g(S1, "A8:D11", "bands_after")
    steps.append({"op": "set", "sheet": S1, "addr": "B17", "value": -86.75})
    steps.append({"op": "set", "sheet": S1, "addr": "B18", "value": 0.25})
    steps.append({"op": "calc", "key": "old_demo"})
    g(S1, "B22:B45", "old_b_polar")
    steps.append({"op": "set", "sheet": S1, "addr": "B17", "value": 0.25})
    steps.append({"op": "calc", "key": "old_demo2"})
    g(S1, "B22:B45", "old_b_eq")
    steps.append({"op": "dv", "sheet": S1B, "addr": "B6", "key": "dv_b6"})
    steps.append({"op": "dv", "sheet": S1B, "addr": "B7", "key": "dv_b7"})
    steps.append({"op": "dv", "sheet": S1C, "addr": "B5", "key": "dv_site"})
    for k, v in (("ok30", 30), ("ok5", 5), ("ok1", 1), ("ng10", 10), ("ng0", 0), ("ngabc", "abc")):
        steps.append({"op": "testvalid", "sheet": S1B, "addr": "B6", "value": v, "key": "v6_" + k})
    for k, v in (("ok50", 50), ("ok90", 90), ("ok85", 85), ("ng49", 49), ("ng91", 91), ("ng85.5", 85.5), ("ngx", "x")):
        steps.append({"op": "testvalid", "sheet": S1B, "addr": "B7", "value": v, "key": "v7_" + k})
    for i, nm in enumerate(SITES):
        steps.append({"op": "testvalid", "sheet": S1C, "addr": "B5", "value": nm, "key": f"v5_ok{i}"})
    steps.append({"op": "testvalid", "sheet": S1C, "addr": "B5", "value": "でたらめ", "key": "v5_ng"})
    steps.append({"op": "testvalid", "sheet": S1C, "addr": "B5", "value": "コペ", "key": "v5_ng2"})
    # 1b：幅 30・5・1 の表（90行）。線は 90（全部内側）
    for wd in (30, 5, 1):
        steps.append({"op": "set", "sheet": S1B, "addr": "B6", "value": wd})
        steps.append({"op": "set", "sheet": S1B, "addr": "B7", "value": 90})
        steps.append({"op": "calc", "key": f"b1_{wd}"})
        g(S1B, "A15:E104", f"b1_tbl_{wd}")
        g(S1B, "B9:B10", f"b1_res_{wd}")
        g(S1B, "A11", f"b1_st_{wd}", True)
    LINES = (50, 70, 80, 85, 86, 87, 88, 90)
    for ln in LINES:
        steps.append({"op": "set", "sheet": S1B, "addr": "B7", "value": ln})
        steps.append({"op": "calc", "key": f"b1_line{ln}"})
        g(S1B, "B9:B10", f"b1_line_{ln}")
        g(S1B, "A11", f"b1_linest_{ln}", True)
    steps.append({"op": "set", "sheet": S1B, "addr": "B7", "value": 95})       # 貼り付けなどで範囲外が入った場合
    steps.append({"op": "calc", "key": "b1_bad7"})
    g(S1B, "B9:B10", "b1_bad7_res", True)
    g(S1B, "A11", "b1_bad7_st", True)
    steps.append({"op": "set", "sheet": S1B, "addr": "B7", "value": 85})
    steps.append({"op": "set", "sheet": S1B, "addr": "B6", "value": 10})
    steps.append({"op": "calc", "key": "b1_bad6"})
    g(S1B, "A11", "b1_bad6_st", True)
    # 1c：5クレーター＋空欄＋でたらめ
    for i, nm in enumerate(SITES):
        steps.append({"op": "set", "sheet": S1C, "addr": "B5", "value": nm})
        steps.append({"op": "calc", "key": f"c1_{i}"})
        g(S1C, "B10:E10", f"c1_row_{i}")
        g(S1C, "B15:C38", f"c1_curve_{i}")
        g(S1C, "A6", f"c1_st_{i}", True)
    for key, val in (("blank", None), ("bad", "でたらめ")):
        steps.append({"op": "set", "sheet": S1C, "addr": "B5", "value": val})
        steps.append({"op": "calc", "key": "c1_" + key})
        g(S1C, "A6", f"c1_{key}_st", True)
        g(S1C, "B10:E10", f"c1_{key}_row", True)
        g(S1C, "B15:C16", f"c1_{key}_curve", True)
    # 見た目の確認用：状態を決めてから、チャートを PNG、シートを PDF に出す（検証用コピー。原本は保存しない）
    steps.append({"op": "set", "sheet": S1B, "addr": "B6", "value": 1})
    steps.append({"op": "set", "sheet": S1B, "addr": "B7", "value": 85})
    steps.append({"op": "set", "sheet": S1C, "addr": "B5", "value": "ラングレヌス"})
    steps.append({"op": "calc", "key": "final"})
    steps.append({"op": "chart", "sheet": S1B, "key": "chart_1b"})
    steps.append({"op": "chart", "sheet": S1C, "key": "chart_1c"})
    steps.append({"op": "exportchart", "sheet": S1B, "path": os.path.join(workdir, "chart_1b"), "key": "png_1b"})
    steps.append({"op": "exportchart", "sheet": S1C, "path": os.path.join(workdir, "chart_1c"), "key": "png_1c"})
    steps.append({"op": "exportpdf", "sheet": S1B, "path": os.path.join(workdir, "sheet_1b.pdf"), "area": "A1:R40", "landscape": True, "fit": True})
    steps.append({"op": "exportpdf", "sheet": S1C, "path": os.path.join(workdir, "sheet_1c.pdf"), "area": "A1:N42", "landscape": True, "fit": True})
    steps.append({"op": "exportpdf", "sheet": S1, "path": os.path.join(workdir, "sheet_1.pdf"), "area": "A1:L58", "fit": True})
    # ステップ4：月全体（線85。温度を1番にした設定A）の状態で PDF に出す（状態表示・上位10・B36 の見た目の目視確認用）
    steps.append({"op": "set", "sheet": S4, "addr": "C5", "value": ALL_MOON})
    steps.append({"op": "set", "sheet": S4, "addr": "C4", "value": 85})
    for j, v in enumerate((1, 3, 0, 0, 0)):
        steps.append({"op": "set", "sheet": S4, "addr": f"C{8 + j}", "value": v})
    steps.append({"op": "calc", "key": "pdf4"})
    steps.append({"op": "exportpdf", "sheet": S4, "path": os.path.join(workdir, "sheet_4.pdf"), "area": "A1:H38", "landscape": True, "fit": True})
    steps.append({"op": "exportpdf", "sheet": "はじめに", "path": os.path.join(workdir, "sheet_intro.pdf"), "area": "A1:B34", "fit": True})

    out = run_excel(steps, path, workdir)
    print(f"   Excel {out.get('excel_version')}：開く {out['open_ms']} ms、初回再計算 {out.get('calc_ms_first')} ms、"
          f"検査全体 {out['total_ms'] / 1000:.0f} 秒")

    errs = out.get("errors") or []
    record(not errs, f"{label}：Excel の操作（値の入力・再計算・読み取り）にエラーなし", "" if not errs else str(errs[:3]))
    # --- 開いて計算できたか／シート ---
    vis = {n: v for n, v, _ in out["info"]["sheets"]}
    want_hidden = set(STUDENT_HIDDEN) if student else set()
    got_hidden = {n for n, v in vis.items() if v != -1}
    record(got_hidden == want_hidden, f"{label}：Excel で開けて、非表示シートが期待どおり", "、".join(sorted(got_hidden)))
    # --- 値 ---
    bands = [int(r[3]) for r in out["bands"]]
    record(bands == ex.bands() == [295, 279, 234, 156], f"{label}：ステップ1の4バンド（Excel）{bands}（期待 295・279・234・156 K）")
    s2 = {r[0]: r[1] for r in out["s2"]}
    ns, nl, ds, dl, r1, r = ex.density()
    vals = [x[1] for x in out["s2"]]
    record(vals[0] == ns and vals[1] == nl and vals[4] == ds and vals[5] == dl and abs(vals[6] - r1) < 1e-9,
           f"{label}：ステップ2（Excel）海{vals[0]}・陸{vals[1]}・密度 {vals[4]}/{vals[5]}・倍率 {vals[6]}（小数1桁表示。厳密な比 {r:.3f}）")
    us, ul = ex.usgs()
    d = out["s2d"]
    record(abs(d[1][1] - us) < 1e-9 and abs(d[2][1] - ul) < 1e-9,
           f"{label}：USGS 年代（Excel）海{round(d[1][1], 2)}・陸{round(d[2][1], 2)}")
    # --- 入力規則 ---
    dv = out["dv_c5"]
    record(dv.get("type") == 3 and dv.get("formula1") == "=AreaChoices" and dv.get("in_cell_dropdown") is True
           and dv.get("alert_style") == 1 and dv.get("show_error") is True,
           f"{label}：C5 のプルダウン（Excel：リスト、数式 {dv.get('formula1')}、停止メッセージ）")
    ok_c5 = all(out[f"c5_ok{i}"] is True for i in range(len(ex.regions)))
    record(ok_c5 and out["c5_ng"] is False and out["c5_ng2"] is False and out["c5_all_ok"] is True and out["c5_all_ng"] is False,
           f"{label}：C5 の入力規則：8 地域名と『{ALL_MOON}』は受理、「でたらめ」「赤道の海」「月全体」（名前の一部）は拒否")
    for lab, suf in (("ステップ4 C8", ""), ("4b C8", "_b")):
        okv = [out[f"w_ok{n}{suf}"] for n in (3, 0, 5)]
        ngv = [out[f"w_ng{n}{suf}"] for n in ("-1", "9", "2.5", "abc")]
        record(all(v is True for v in okv) and all(v is False for v in ngv),
               f"{label}：重み（{lab}）の入力規則：0・3・5 は受理、-1・9・2.5・abc は拒否")
    wdv = out["dv_w4"]
    record(wdv.get("type") == 1 and wdv.get("formula1") == "0" and wdv.get("formula2") == "5" and wdv.get("operator") == 1,
           f"{label}：重みの入力規則（Excel：整数、0〜5、{wdv.get('formula1')}〜{wdv.get('formula2')}）")
    # --- 上位10：pandas と一致、10行が全部異なる ---
    bad_dist, bad_eq = [], []
    n_dist = 0
    for key, kind, reg, w in cases:
        if kind == 4:
            rows, pts, sc = ex.top4(reg, w)
            want = min(10, int((ex.env["region"] == reg).sum())) if sum(w) > 0 else 0
        else:
            if sum(w) == 0:
                rows, pts = [], []
            else:
                rows, pts, sc = ex.top4b(w)
            want = 10 if sum(w) > 0 else 0
        got = top_from(out[key + "_top"])
        if kind == "4b" and key == "s4b_over":
            continue  # 重み 9 は入力規則の外。状態表示の検査で扱う
        if len(set(got)) != len(got) or len(got) != want:
            bad_dist.append((key, len(got), len(set(got))))
        if got != pts:
            bad_eq.append(key)
        if key.startswith("s4_") or key.startswith("s4b_single"):
            n_dist += 1
    record(not bad_dist, f"{label}：Excel の上位10が全部異なる地点（全地域×単一指標 40＋4b 単一指標 4＋式の検査例 ほか）",
           "" if not bad_dist else str(bad_dist[:5]))
    record(not bad_eq, f"{label}：Excel の上位10（緯度・経度、全 {len(cases) - 1} 通り）が pandas の再現と一致",
           "" if not bad_eq else str(bad_eq[:8]))
    # 式の検査例の1位
    for lab, (reg, w) in STD4.items():
        got = top_from(out[f"std4_{lab}_top"])
        ro, po, _ = ex.top4(reg, list(w), "old")
        record(got[0] == po[0], f"{label}：{lab} の1位（Excel）{got[0]}は変更前の1位 {po[0]} と同じ")
    for lab, w in STD4B.items():
        got = top_from(out[f"std4b_{lab}_top"])
        ro, po, _ = ex.top4b(list(w), "old")
        record(got[0] == po[0], f"{label}：{lab} の1位（Excel）{got[0]}は変更前の1位 {po[0]} と同じ")
    # --- 表示のスコアの見た目 ---
    sc_rows = out[f"std4_{K_STD4_FAR}_score"]
    reg, w = STD4[K_STD4_FAR]
    _, _, sc = ex.top4(reg, list(w))
    rows, _, _ = ex.top4(reg, list(w))
    exp_scores = [f"{round(sc[x - 2], 3):.3f}".rstrip("0").rstrip(".") for x in rows]
    got_scores = [r[0] for r in sc_rows]
    ok_sc = all(abs(float(a) - round(sc[x - 2], 3)) < 1e-9 for a, x in zip(got_scores, rows))
    record(ok_sc, f"{label}：表示のスコアは ROUND(元のスコア,3) と一致（見た目は変わらない）", "/".join(got_scores[:5]) + "…")
    # --- A4 範囲表示 ---
    for lab, (reg, w) in STD4.items():
        rows, pts, _ = ex.top4(reg, list(w))
        txt = out[f"std4_{lab}_rng"][0][0]
        record(txt == range_text4(pts), f"{label}：A4 範囲表示 {lab}", txt)
    for lab, w in STD4B.items():
        rows, pts, _ = ex.top4b(list(w))
        txt = out[f"std4b_{lab}_rng"][0][0]
        record(txt == range_text4b(ex, rows), f"{label}：A4 範囲表示 {lab}", txt)
    # 日付変更線をまたぐ領域（南極エイトケン・南極）は経度の範囲を出さない
    key = "s4_%d_%d" % (ex.regions.index("裏側・南極エイトケン（フォン・カルマン）"), 0)
    # --- 状態表示 ---
    st = lambda k: out[k + "_st"][0][0]
    record(st("c5_blank").startswith("地域タイプを選んでください") and all(r[0] == "" for r in out["c5_blank_lat"]),
           f"{label}：C5 が空欄 → 「{st('c5_blank')}」、上位10は空白")
    record(st("c5_bad").startswith("【注意】C5 の地域名が") and all(r[0] == "" for r in out["c5_bad_lat"]),
           f"{label}：C5 が不一致 → 「{st('c5_bad')}」、上位10は空白")
    record(st("s4_zero").startswith("【注意】重みがすべて0"), f"{label}：ステップ4 重みが全部0 → 「{st('s4_zero')}」")
    record(st("s4b_zero").startswith("【注意】重みがすべて0"), f"{label}：4b 重みが全部0 → 「{st('s4b_zero')}」")
    record(st("s4b_over").startswith("【注意】重みは0〜5の整数"), f"{label}：4b 重みが範囲外（貼り付けで 9）→ 「{st('s4b_over')}」")
    ok_regs = []
    for ri, nm in enumerate(ex.regions):
        t = out[f"reg{ri}_st"][0][0]
        n = int((ex.env["region"] == nm).sum())
        if nm.startswith("南極"):
            ok_regs.append(t.startswith("南極は『ステップ4b_スコア』を使います"))
        else:
            ok_regs.append(t.startswith(f"OK　採点対象 {n} 行"))
    record(all(ok_regs), f"{label}：8 地域を C5 に入れたときの状態表示（南極は「南極は『ステップ4b_スコア』を使います…」、他は OK＋行数）",
           f"南極：{out['reg6_st'][0][0]}")
    # 日付変更線
    key = f"s4_{ex.regions.index('裏側・南極エイトケン（フォン・カルマン）')}_3"
    rng = out[key + "_rng"][0][0]
    rows, pts, _ = ex.top4("裏側・南極エイトケン（フォン・カルマン）", [0, 0, 0, 3, 0])
    record(rng == range_text4(pts), f"{label}：日付変更線をまたぐ領域の範囲表示：{rng}")

    check_excel_dai2ji(ex, out, label, student, LINES_CHECK)
    check_excel_v3(ex, out, label, student)
    # --- G10 チャートの位置 ---
    ch = out["chart1"][0]
    record(ch[1] == "$D$21", f"{label}：ステップ1のチャートは D21 起点（注意書き A19:H19 に重ならない）", f"{ch[1]}〜{ch[2]}")
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--teacher", default=str(TEACHER))
    ap.add_argument("--student", default=str(STUDENT))
    ap.add_argument("--pandas", action="store_true")
    ap.add_argument("--excel", action="store_true")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--workdir", default=None, help="Excel に渡すコピーと一時ファイルの置き場所")
    a = ap.parse_args()
    if a.all:
        a.pandas = a.excel = True

    check_answers(a.teacher, a.student)
    check_mission_words(a.teacher, a.student)
    check_v3_student_words(a.teacher, a.student)
    check_region_names(a.teacher, "教員版")
    check_region_names(a.student, "学習者版")
    check_pii(a.teacher, "教員版")
    check_pii(a.student, "学習者版")
    check_structure(a.teacher, a.student)
    ex = None
    if a.pandas or a.excel:
        ex = Expect()
    if a.pandas:
        check_pandas(ex)
    if a.excel:
        work = a.workdir or tempfile.mkdtemp(prefix="course_xlsx_check_")
        for label, path, student in (("教員版", a.teacher, False), ("学習者版", a.student, True)):
            check_excel(ex, path, label, os.path.join(work, "t" if not student else "s"), student)

    ng = [r for r in RESULTS if not r[0]]
    print(f"\n合計 {len(RESULTS)} 項目：合格 {len(RESULTS) - len(ng)}、不合格 {len(ng)}")
    for _, n, d in ng:
        print("  不合格：", n, d)
    return 1 if ng else 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
